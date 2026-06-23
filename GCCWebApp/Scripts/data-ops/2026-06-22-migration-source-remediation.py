#!/usr/bin/env python3
"""
Migration source-data MECHANICAL remediation  (DEV / TEST ONLY)
================================================================

Applies ONLY the deterministic, decision-free fixes from
docs/tracking/migration-remediation-plan-2026-06-22.md to STAGING COPIES of the
customer source workbooks. Originals are never modified.

What it fixes (mechanical):
  ACT-CAT     fix 5 mis-numbered product-category prefixes (label is canonical)
  EXP-CERT-DUP drop exact-duplicate rows in the Expired list (keep first)
  EXP-DATES   coerce Excel date serials (e.g. 45536) -> real dates (Active+Expired)
  ASIN-JUNK   exclude rows whose product-id value is not a valid ASIN (placeholders)
  ASIN-DUP    drop exact-duplicate ASIN rows (keep first)

What it deliberately does NOT do (left for decisions / import-time transform):
  - widen productName, results tie-break, company canonicalization        (client decision)
  - retired cert types, orphan ASINs                                      (client decision, 06-11 note)
  - authoritative-source collapse for cert-level ASIN-CONFLICT fields     (import transform)
  - blank/N/A category backfill (ACT-CAT-MISSING)                         (analyst pass)
  - residual same-cert# differing rows, blank cert#, exp<eff              (reported only)

Outputs (under ./_remediation_output/<timestamp>/):
  cleaned_cert_listing.xlsx   - all sheets preserved; Active+Expired remediated
  cleaned_asin_list.xlsx      - single sheet, junk + exact dups removed
  remediation-change-log.csv  - every cell changed and row dropped (before/after)
  residual-issues.csv         - decision-dependent items left untouched

Usage:
  python 2026-06-22-migration-source-remediation.py
  python 2026-06-22-migration-source-remediation.py <cert.xlsx> <asin.xlsm>

Then re-run the gate against the cleaned files:
  python 2026-06-22-migration-source-validation.py \
      _remediation_output/<ts>/cleaned_cert_listing.xlsx \
      _remediation_output/<ts>/cleaned_asin_list.xlsx

Requires: openpyxl
"""

from __future__ import annotations

import csv
import datetime
import os
import re
import shutil
import sys
import warnings
from datetime import timedelta

warnings.filterwarnings("ignore")

try:
    import openpyxl
except ImportError:
    sys.exit("ERROR: openpyxl is required.  Run:  pip install openpyxl")

# --------------------------------------------------------------------------- #
# Config (mirrors the validation script)
# --------------------------------------------------------------------------- #

DEFAULT_DIR = (
    r"C:\Users\swest\OneDrive - PointSolve Technology Inc\LocalWest"
    r"\Projects\GCC Database Solution\TestData\DBSampleLists"
)
DEFAULT_CERT_FILE = "GCC Database SAMPLE Cert Listing for Anuki Tesing_6.2.26.xlsx"
DEFAULT_ASIN_FILE = "SAMPLE ASIN List for Anuki Testing_6.1.26.xlsm"

# Canonical Level-1 category label -> fixed numeric prefix (Field Options sheet).
CANONICAL_CATEGORIES = {
    "Chemicals and Raw Materials": "1",
    "Clothing, Apparel, and Fabrics": "2",
    "Construction / Building Products": "3",
    "Consumer Products": "4",
    "Electronics, Computers, and Accessories": "5",
    "Food and Beverage": "6",
    "Healthcare and Pharmaceutical": "7",
    "Packaging": "8",
}

ASIN_PATTERN = re.compile(r"B0[A-Z0-9]{8}$")
EXCEL_EPOCH = datetime.datetime(1899, 12, 30)  # Windows 1900 date system
SERIAL_MIN, SERIAL_MAX = 30000, 60000  # ~1982..2064, the plausible date-serial band

# --------------------------------------------------------------------------- #
# Change log
# --------------------------------------------------------------------------- #

changes = []   # action, file, sheet, excel_row, cert#, field, before, after, reason
residuals = []  # file, sheet, excel_row, cert#, issue, detail


def log_change(action, file, sheet, row, cert, field, before, after, reason):
    changes.append({"action": action, "file": file, "sheet": sheet, "excel_row": row,
                    "cert#": cert, "field": field, "before": before, "after": after,
                    "reason": reason})


def log_residual(file, sheet, row, cert, issue, detail):
    residuals.append({"file": file, "sheet": sheet, "excel_row": row, "cert#": cert,
                      "issue": issue, "detail": detail})


def clean(v):
    return "" if v is None else str(v).strip()


def coerce_serial(v):
    """Return (datetime, True) if v is an Excel date serial, else (v, False)."""
    if isinstance(v, datetime.datetime):
        return v, False
    if isinstance(v, bool):
        return v, False
    if isinstance(v, (int, float)) and SERIAL_MIN <= v <= SERIAL_MAX:
        return EXCEL_EPOCH + timedelta(days=int(v)), True
    s = clean(v)
    if s.isdigit() and SERIAL_MIN <= int(s) <= SERIAL_MAX:
        return EXCEL_EPOCH + timedelta(days=int(s)), True
    return v, False


# --------------------------------------------------------------------------- #
# Cert workbook (edit in place on a copy -- preserves all 4 sheets)
# --------------------------------------------------------------------------- #


def find_sheet(wb, predicate):
    for ws in wb.worksheets:
        if predicate(ws.title):
            return ws
    return None


def remediate_cert(src_path, out_path):
    shutil.copy2(src_path, out_path)
    wb = openpyxl.load_workbook(out_path)  # not read_only -> editable; keeps formulas/sheets
    fname = "cleaned_cert_listing.xlsx"

    active = find_sheet(wb, lambda t: t.strip().lower().startswith("active list"))
    expired = find_sheet(wb, lambda t: t.strip().lower().startswith("expired list"))

    # ---- Active: category prefix fix + date coercion ----
    # cols (1-based): cert#=1, category=4, effective=8, expiration=9
    A_CERT, A_CAT, A_EFF, A_EXP = 1, 4, 8, 9
    for row in range(2, active.max_row + 1):
        cert = clean(active.cell(row, A_CERT).value)
        if not cert and not any(clean(active.cell(row, c).value) for c in range(1, active.max_column + 1)):
            continue  # blank spacer row

        # category prefix
        cell = active.cell(row, A_CAT)
        raw = clean(cell.value)
        m = re.match(r"^(\d+)\s+(.*)$", raw)
        if m:
            num, label = m.group(1), m.group(2).strip()
            exp = CANONICAL_CATEGORIES.get(label)
            if exp and exp != num:
                new = f"{exp} {label}"
                cell.value = new
                log_change("CELL-FIX", fname, active.title, row, cert, "Product Category",
                           raw, new, "ACT-CAT: prefix corrected to canonical for label")
            elif exp is None:
                log_residual(fname, active.title, row, cert, "ACT-CAT-UNKNOWN",
                             f"category label not canonical: {raw!r}")
        elif raw and raw.upper() != "N/A":
            log_residual(fname, active.title, row, cert, "ACT-CAT-NOPREFIX",
                         f"category has no numeric prefix: {raw!r}")
        elif raw == "" or raw.upper() == "N/A":
            log_residual(fname, active.title, row, cert, "ACT-CAT-MISSING",
                         f"blank/N-A category (analyst backfill): {raw!r}")

        # dates
        for c, field in ((A_EFF, "EffectiveDate"), (A_EXP, "ExpirationDate")):
            cell = active.cell(row, c)
            new, did = coerce_serial(cell.value)
            if did:
                cell.value = new
                cell.number_format = "yyyy-mm-dd"
                log_change("CELL-FIX", fname, active.title, row, cert, field,
                           "<serial>", new.date().isoformat(), "EXP-DATES: Excel serial -> date")
            elif clean(cell.value) and not isinstance(cell.value, datetime.datetime):
                log_residual(fname, active.title, row, cert, "BAD-DATE",
                             f"{field} not a date: {clean(cell.value)!r}")
        eff, exp = active.cell(row, A_EFF).value, active.cell(row, A_EXP).value
        if isinstance(eff, datetime.datetime) and isinstance(exp, datetime.datetime) and exp < eff:
            log_residual(fname, active.title, row, cert, "DATE-RANGE",
                         f"expiration {exp.date()} < effective {eff.date()}")

    # ---- Expired: date coercion + exact-duplicate row removal ----
    # cols (1-based): cert#=1, effective=6, expiration=7
    E_CERT, E_EFF, E_EXP = 1, 6, 7
    maxc = expired.max_column
    # first pass: date coercion + residual reporting
    for row in range(2, expired.max_row + 1):
        if not any(clean(expired.cell(row, c).value) for c in range(1, maxc + 1)):
            continue
        cert = clean(expired.cell(row, E_CERT).value)
        for c, field in ((E_EFF, "EffectiveDate"), (E_EXP, "ExpirationDate")):
            cell = expired.cell(row, c)
            new, did = coerce_serial(cell.value)
            if did:
                cell.value = new
                cell.number_format = "yyyy-mm-dd"
                log_change("CELL-FIX", fname, expired.title, row, cert, field,
                           "<serial>", new.date().isoformat(), "EXP-DATES: Excel serial -> date")
        if not cert:
            log_residual(fname, expired.title, row, cert, "EXP-CERT-BLANK",
                         "expired row has no certificate number (decision: supply or drop)")

    # second pass: exact-duplicate detection (whole row) + residual same-cert# differing rows
    seen_rows = {}      # row-tuple -> first excel row
    cert_rows = {}      # cert# -> list of (excel_row, row-tuple)
    delete_rows = []
    for row in range(2, expired.max_row + 1):
        vals = tuple(clean(expired.cell(row, c).value) for c in range(1, maxc + 1))
        if not any(vals):
            continue
        cert = vals[0]
        if vals in seen_rows:
            delete_rows.append(row)
            log_change("ROW-DROP", fname, expired.title, row, cert, "(entire row)",
                       f"dup of row {seen_rows[vals]}", "", "EXP-CERT-DUP: exact-duplicate row removed")
        else:
            seen_rows[vals] = row
            cert_rows.setdefault(cert, []).append((row, vals))
    # residual: same cert# surviving on >1 DIFFERING rows
    for cert, lst in cert_rows.items():
        if cert and len(lst) > 1:
            log_residual(fname, expired.title, lst[0][0], cert, "EXP-CERT-DUP-RESIDUAL",
                         f"cert# on {len(lst)} differing rows (decision: which is authoritative) "
                         f"rows={[r for r, _ in lst]}")
    for row in sorted(delete_rows, reverse=True):
        expired.delete_rows(row, 1)

    wb.save(out_path)
    return len(delete_rows)


# --------------------------------------------------------------------------- #
# ASIN workbook (write fresh single-sheet xlsx)
# --------------------------------------------------------------------------- #


def remediate_asin(src_path, out_path):
    fname = "cleaned_asin_list.xlsx"
    wb_in = openpyxl.load_workbook(src_path, read_only=True, data_only=True)
    ws_in = None
    for ws in wb_in.worksheets:
        if "upload template" in ws.title.strip().lower():
            ws_in = ws
            break
    if ws_in is None:
        wb_in.close()
        raise KeyError("No 'upload template' sheet in ASIN file")

    it = ws_in.iter_rows(values_only=True)
    header = list(next(it))
    C_TYPE, C_VALUE, C_CERT = 1, 2, 7

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = "GreenCircle Upload Template"
    out_ws.append(header)

    seen = {}
    excel_row = 1  # header is row 1
    junk_dropped = dup_dropped = kept = 0
    for r in it:
        excel_row += 1
        if not any(c is not None and str(c).strip() for c in r):
            continue
        cert = clean(r[C_CERT]) if C_CERT < len(r) else ""
        value = clean(r[C_VALUE]) if C_VALUE < len(r) else ""
        idtype = clean(r[C_TYPE]) if C_TYPE < len(r) else ""

        # ASIN-JUNK: exclude non-ASIN placeholder values
        if idtype == "ASIN" and not ASIN_PATTERN.match(value):
            junk_dropped += 1
            log_change("ROW-EXCLUDE", fname, out_ws.title, excel_row, cert, "value",
                       value, "(excluded)", "ASIN-JUNK: not a valid ASIN")
            continue

        # ASIN-DUP: drop exact-duplicate rows (first 16 cols)
        key = tuple(clean(c) for c in r[:16])
        if key in seen:
            dup_dropped += 1
            log_change("ROW-DROP", fname, out_ws.title, excel_row, cert, "(entire row)",
                       f"dup of row {seen[key]}", "", "ASIN-DUP: exact-duplicate row removed")
            continue
        seen[key] = excel_row
        out_ws.append(list(r))
        kept += 1

    wb_in.close()
    out_wb.save(out_path)
    return junk_dropped, dup_dropped, kept


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #


def write_csv(path, rows, fieldnames):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)


def main():
    argv = sys.argv[1:]
    if len(argv) >= 2:
        cert_src, asin_src = argv[0], argv[1]
    else:
        cert_src = os.path.join(DEFAULT_DIR, DEFAULT_CERT_FILE)
        asin_src = os.path.join(DEFAULT_DIR, DEFAULT_ASIN_FILE)
    for p in (cert_src, asin_src):
        if not os.path.exists(p):
            sys.exit(f"ERROR: file not found: {p}")

    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    outdir = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          "_remediation_output", stamp)
    os.makedirs(outdir, exist_ok=True)

    cert_out = os.path.join(outdir, "cleaned_cert_listing.xlsx")
    asin_out = os.path.join(outdir, "cleaned_asin_list.xlsx")

    print(f"Source cert: {cert_src}")
    print(f"Source ASIN: {asin_src}")
    print("Applying mechanical fixes (originals untouched)...\n")

    exp_dropped = remediate_cert(cert_src, cert_out)
    junk, asin_dup, asin_kept = remediate_asin(asin_src, asin_out)

    write_csv(os.path.join(outdir, "remediation-change-log.csv"), changes,
              ["action", "file", "sheet", "excel_row", "cert#", "field", "before", "after", "reason"])
    write_csv(os.path.join(outdir, "residual-issues.csv"), residuals,
              ["file", "sheet", "excel_row", "cert#", "issue", "detail"])

    cell_fixes = sum(1 for c in changes if c["action"] == "CELL-FIX")
    print("MECHANICAL FIXES APPLIED")
    print("-" * 60)
    print(f"  Cell fixes (category + dates) : {cell_fixes}")
    print(f"  Expired exact-dup rows dropped: {exp_dropped}")
    print(f"  ASIN junk rows excluded       : {junk}")
    print(f"  ASIN exact-dup rows dropped   : {asin_dup}")
    print(f"  ASIN rows kept                : {asin_kept}")
    print(f"\n  Residual items (need decision/analyst): {len(residuals)}")
    from collections import Counter
    for issue, n in sorted(Counter(r["issue"] for r in residuals).items()):
        print(f"     {issue:22s} {n}")
    print(f"\nOutputs written to:\n  {outdir}")
    print("  - cleaned_cert_listing.xlsx")
    print("  - cleaned_asin_list.xlsx")
    print("  - remediation-change-log.csv")
    print("  - residual-issues.csv")
    print("\nNext: re-run the validation gate against the cleaned files to confirm "
          "the mechanical blockers are cleared.")


if __name__ == "__main__":
    main()
