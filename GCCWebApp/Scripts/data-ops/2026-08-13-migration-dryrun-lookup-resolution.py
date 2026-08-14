#!/usr/bin/env python3
"""
Migration DRY RUN — lookup resolution gate            (READ-ONLY; writes no DB rows)
===================================================================================

Answers one question the go-live sequence depends on:

    If we purge the test data but PRESERVE the reference tables the client named
    (Certification Types, Product Category/Type/Subcategory, Tags, Pathways,
    Partners), will the 5,556-certificate source load actually resolve against
    what is left?

Every source value that does NOT resolve is either (a) a row that will fail on
import, or (b) worse, a row that imports with the field silently blank. This
script finds them before anything is deleted or loaded.

Inputs
------
  <cert.xlsx>   the client cert workbook (default: DataMigrationTestDev\\GCC Database Migration_Certs_8.12.26V2.xlsx)
  <lookupdir>   directory of pipe-delimited lookup dumps taken from the TARGET db:
                  lk_certtypes.txt  Id|Name|Level|ParentId|IsActive
                  lk_prodcat.txt    Id|Name
                  lk_prodtype.txt   Id|Name
                  lk_prodsub.txt    Id|Name
                  lk_pathways.txt   Id|Name
                  lk_partners.txt   Id|Name
                  lk_tags.txt       Id|Name|TagType

Transform rules applied before matching (per plan §11.3)
--------------------------------------------------------
  * Numeric grouping prefixes are stripped: "4.3 Cleaning Supplies" -> "Cleaning Supplies".
    The prefix is a spreadsheet grouping aid and must not reach the DB.
  * Matching is case-insensitive and whitespace-collapsed, mirroring the importer's
    case-insensitive lookup resolution.
  * "N/A" is reported as its own bucket, NOT as unresolved — the client states it is a
    legitimate option for certs that are not product-level.

Outputs (under ./_dryrun_output/<timestamp>/)
  lookup-resolution-summary.csv    one row per (sheet, field): resolved / unresolved / blank
  lookup-unresolved-values.csv     every distinct value that does not resolve, with row counts

Usage
  python 2026-08-13-migration-dryrun-lookup-resolution.py
  python 2026-08-13-migration-dryrun-lookup-resolution.py <cert.xlsx> <lookupdir>

Requires: openpyxl
"""

from __future__ import annotations

import csv
import datetime
import os
import re
import sys
from collections import Counter, defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))

DEFAULT_CERTS = os.path.join(
    REPO_ROOT, "DataMigrationTestDev", "GCC Database Migration_Certs_8.12.26V2.xlsx"
)

# Sheets that are actually load input (plan §11.1). Full/Hold lists are context, not input.
LOAD_SHEETS = [
    "Active List 7.31.26",
    "Expired List 2 - 8.11.26",
    "Expired List 1 - 4.9.26",
]

# source column header -> lookup file it must resolve against.
# Expired List 1 carries only CERTIFICATE_TYPE; the others are absent there by design.
FIELD_MAP = [
    ("CERTIFICATE_TYPE", "certtypes"),
    ("Product Category", "prodcat"),
    ("Product Type", "prodtype"),
    ("Product Subcategory", "prodsub"),
    ("Pathway (SEP Amazon only)", "pathways"),
    ("Partner (Amazon)", "partners"),
    ("External Tags", "tags"),
    ("Internal Tags", "tags"),
]

# "4.3 Cleaning Supplies" / "12 Foo" / "4.3.1 Bar" -> strip the leading numeric group.
PREFIX_RE = re.compile(r"^\d+(?:\.\d+)*\s+")
NA_VALUES = {"n/a", "na", "n\\a", "none"}


def norm(v) -> str:
    """Case/whitespace-insensitive match key, after stripping the grouping prefix."""
    if v is None:
        return ""
    s = str(v).replace(" ", " ").strip()
    s = PREFIX_RE.sub("", s)
    return re.sub(r"\s+", " ", s).strip().casefold()


def strip_prefix(v) -> str:
    if v is None:
        return ""
    s = str(v).replace(" ", " ").strip()
    return re.sub(r"\s+", " ", PREFIX_RE.sub("", s)).strip()


def read_lookup(path: str):
    """Read a sqlcmd pipe dump -> (set of normalized names, list of raw names)."""
    keys, raw = set(), []
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        return keys, raw
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        for line in fh:
            line = line.strip().lstrip("﻿")
            if not line or line.startswith("-") or "rows affected" in line:
                continue
            parts = [p.strip() for p in line.split("|")]
            if len(parts) < 2:
                continue
            name = parts[1]
            if not name or name.upper() == "NULL":
                continue
            raw.append(name)
            keys.add(norm(name))
    return keys, raw


def header_index(ws):
    """header text (trimmed) -> list of column indexes. Duplicated headers are kept."""
    idx = defaultdict(list)
    for row in ws.iter_rows(min_row=1, max_row=1):
        for i, cell in enumerate(row):
            if cell.value:
                idx[str(cell.value).strip()].append(i)
    return idx


def resolve_header(idx, wanted):
    """Match a header tolerantly — the source has trailing spaces ('Product Type ')
    and long parenthetical suffixes that differ per sheet."""
    w = wanted.strip().casefold()
    for h, cols in idx.items():
        if h.strip().casefold() == w:
            return cols
    for h, cols in idx.items():
        if h.strip().casefold().startswith(w):
            return cols
    return []


def main():
    certs_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_CERTS
    lookup_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "_dryrun_lookups")

    if not os.path.exists(certs_path):
        sys.exit(f"Cert workbook not found: {certs_path}")
    if not os.path.isdir(lookup_dir):
        sys.exit(f"Lookup dump dir not found: {lookup_dir}")

    lookups = {}
    for key in ("certtypes", "prodcat", "prodtype", "prodsub", "pathways", "partners", "tags"):
        keys, raw = read_lookup(os.path.join(lookup_dir, f"lk_{key}.txt"))
        lookups[key] = keys
        print(f"lookup {key:10} {len(keys):4} distinct names in target DB")
    print()

    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    outdir = os.path.join(HERE, "_dryrun_output", stamp)
    os.makedirs(outdir, exist_ok=True)

    wb = openpyxl.load_workbook(certs_path, read_only=True, data_only=True)

    summary_rows = []
    unresolved_rows = []

    for sheet in LOAD_SHEETS:
        if sheet not in wb.sheetnames:
            print(f"!! sheet missing: {sheet}")
            continue
        ws = wb[sheet]
        idx = header_index(ws)

        # collect the column indexes we care about once, then stream the rows once.
        wanted = []
        for header, lk in FIELD_MAP:
            cols = resolve_header(idx, header)
            if cols:
                wanted.append((header, lk, cols))

        counters = {(h, lk): Counter() for h, lk, _ in wanted}
        blanks = {(h, lk): 0 for h, lk, _ in wanted}
        na_counts = {(h, lk): 0 for h, lk, _ in wanted}
        total = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            total += 1
            for header, lk, cols in wanted:
                for ci in cols:
                    val = row[ci] if ci < len(row) else None
                    s = strip_prefix(val)
                    if not s:
                        # only count blank once per field, not once per duplicate column
                        if ci == cols[0]:
                            blanks[(header, lk)] += 1
                        continue
                    if norm(s) in NA_VALUES:
                        na_counts[(header, lk)] += 1
                        continue
                    counters[(header, lk)][s] += 1

        print(f"== {sheet}   ({total} data rows)")
        for header, lk, _cols in wanted:
            c = counters[(header, lk)]
            target = lookups.get(lk, set())
            ok_rows = ok_vals = bad_rows = bad_vals = 0
            for val, n in c.items():
                if norm(val) in target:
                    ok_rows += n
                    ok_vals += 1
                else:
                    bad_rows += n
                    bad_vals += 1
                    unresolved_rows.append(
                        {
                            "Sheet": sheet,
                            "Field": header,
                            "LookupTable": lk,
                            "UnresolvedValue": val,
                            "RowCount": n,
                        }
                    )
            flag = "OK " if bad_vals == 0 else "FAIL"
            print(
                f"   {flag} {header[:44]:44} resolved {ok_vals:3} vals/{ok_rows:5} rows | "
                f"UNRESOLVED {bad_vals:3} vals/{bad_rows:5} rows | "
                f"blank {blanks[(header, lk)]:5} | N/A {na_counts[(header, lk)]:5}"
            )
            summary_rows.append(
                {
                    "Sheet": sheet,
                    "Field": header,
                    "LookupTable": lk,
                    "DataRows": total,
                    "ResolvedDistinct": ok_vals,
                    "ResolvedRows": ok_rows,
                    "UnresolvedDistinct": bad_vals,
                    "UnresolvedRows": bad_rows,
                    "BlankRows": blanks[(header, lk)],
                    "NaRows": na_counts[(header, lk)],
                }
            )
        print()

    with open(os.path.join(outdir, "lookup-resolution-summary.csv"), "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(summary_rows[0].keys()))
        w.writeheader()
        w.writerows(summary_rows)

    with open(os.path.join(outdir, "lookup-unresolved-values.csv"), "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["Sheet", "Field", "LookupTable", "UnresolvedValue", "RowCount"])
        w.writeheader()
        w.writerows(sorted(unresolved_rows, key=lambda r: (-r["RowCount"], r["Sheet"], r["Field"])))

    tot_bad = sum(r["UnresolvedRows"] for r in summary_rows)
    print(f"TOTAL unresolved field-values: {len(unresolved_rows)} distinct across {tot_bad} row-hits")
    print(f"Output: {outdir}")


if __name__ == "__main__":
    main()
