#!/usr/bin/env python3
"""
MIGRATION TRANSFORM  —  client source workbooks -> importer-ready CSV
=====================================================================
Transforms A + B (cert listings -> certificate import CSV) and C (ASIN master
list -> per-certificate product-ID CSV). Read-only against the source files;
writes only into ./_transform_output/<timestamp>/.

This is the step plan §12.5 identified as missing. It does NOT load anything.

------------------------------------------------------------------------------
HEADERS ARE EMITTED EXACTLY AS THE PARSER EXPECTS — DO NOT "TIDY" THEM
------------------------------------------------------------------------------
CsvHelper header matching in CertificateImportService is **case-SENSITIVE**
(no PrepareHeaderForMatch is configured, :1070). That is the DEF-347 failure
mode: a file headed "COUNTRY" silently imported a blank country and raised no
error. Every header below is copied from the column arrays at
CertificateImportService.cs:69-119 and verified to match a listed form.

------------------------------------------------------------------------------
TWO SOURCE FIELDS HAVE NO IMPORT PATH  (DEF-348)
------------------------------------------------------------------------------
The downloaded template (CertificateImportController.cs:464) offers columns
"Cert Subcategories" and "Subcategory", but CertificateImportService has **no
mapping array for either** — ProcessRowAsync passes null for subcategory
outright (:771, "Subcategory still resolves as null (no column yet)").

So those two columns are accepted by the file format and then silently dropped.
This script still emits them (harmless, and correct the day the importer gains
them) but ALSO writes every value that will not land to
`unmappable-fields.csv`, so the loss is counted rather than discovered later.
Do not treat the cert CSV as complete for these two fields.

------------------------------------------------------------------------------
DUPLICATE CERTIFICATE NUMBERS  (plan §11.2 — 175 rows, DECISION STILL OPEN)
------------------------------------------------------------------------------
Certificate.Number carries a UNIQUE index. --dup-policy controls the handling
and DEFAULTS TO THE NON-DESTRUCTIVE OPTION:

  quarantine (default)  every row of an affected number goes to
                        quarantine-duplicate-cert-numbers.csv and NONE of them
                        are emitted for load. Nothing is silently lost and
                        nothing is silently collapsed.
  collapse              keep the most-complete row per number (plan's June
                        recommendation, (a)); the rest still go to quarantine
                        so the drop is auditable. LOSSY — the Regupol rows are
                        genuinely different products (§11.2).
  all                   emit every row. The load WILL fail on the unique index.
                        Provided only to measure the failure, never to run.

Usage
  python 2026-08-13-migration-transform.py
  python 2026-08-13-migration-transform.py --dup-policy collapse
  python 2026-08-13-migration-transform.py <cert.xlsx> <asin.xlsm> --dup-policy quarantine

Requires: openpyxl
"""

from __future__ import annotations

import argparse
import csv
import datetime
import os
import re
import sys
from collections import Counter, defaultdict

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))
SRC_DIR = os.path.join(REPO_ROOT, "DataMigrationTestDev")

DEFAULT_CERTS = os.path.join(SRC_DIR, "GCC Database Migration_Certs_8.12.26V2.xlsx")
DEFAULT_ASINS = os.path.join(SRC_DIR, "GCC ASIN Master List_8.12.26.xlsm")

# Output headers. Each is verified against CertificateImportService.cs:69-119.
# "Cert Subcategories" and "Subcategory" are emitted but NOT ingested — see DEF-348.
CERT_HEADERS = [
    "Certificate #",        # CertificateNumberColumns
    "Certification Type",   # CertificationTypeColumns
    "Cert Subcategories",   # NO PARSER MAPPING - DEF-348
    "Status",               # StatusColumns
    "Company",              # CompanyNameColumns
    "Company Email",        # CompanyEmailColumns
    "Company URL",          # CompanyWebsiteColumns
    "Product Name",         # ProductNameColumns
    "Product Category",     # ProductCategoryColumns
    "Product Type",         # ProductTypeColumns
    "Subcategory",          # NO PARSER MAPPING - DEF-348
    "Partner",              # PartnerColumns
    "Pathway",              # PathwayColumns
    "Cert Results",         # CertificationResultsColumns
    "Effective Date",       # EffectiveDateColumns
    "Expiration Date",      # ExpirationDateColumns
    "Location",             # CertificateLocationColumns
    "Country",              # CertificateCountryColumns
    "External Tag",         # TagsColumns
    "Internal Tag",         # InternalTagsColumns
]

# Product-ID channel headers.
#
# "Partner" is REQUIRED and is not cosmetic. CR-171 Model A
# (Controllers/Api/v1/CertificatesController.cs:270) returns a product ID to a
# partner caller when `!pi.Partners.Any() || pi.Partners.Any(...)` — i.e. an
# UNTAGGED product ID is treated as SHARED and is visible to EVERY partner.
# Omitting the column does not hide the ASINs, it over-exposes them.
#
# "amazonCode" is deliberately NOT emitted: ProductIdService's CSV parser has no
# mapping for it and CertificateProductIds has no such column. The Amazon API
# sources amazonCode from CertificationType.AmazonCode instead
# (CertificatesController.cs:265), which is populated — so the sheet's column is
# redundant, not lost.
#
# ⚠️ DEF-351: DO NOT FEED THESE CSVs TO /api/productids/certificate/{id}/import.
# That parser does a naive `line.Split(',')` with no quote handling, so any
# productName containing a comma (50.7% of this data) is silently truncated at
# the first one. Load via the JSON endpoint instead:
#     POST /api/productids/certificate/{id}/bulk
#     { "Items": [ { Type, Value, Partner, ProductName } ], "OverwriteExisting": false }
# These CSVs remain the reviewable artefact and the input to that JSON post.
PRODUCTID_HEADERS = ["type", "value", "Partner", "productName"]

# PARTNER TAG IS DERIVED FROM THE CERTIFICATE, NOT FROM THE IDENTIFIER TYPE.
#
# An earlier revision hardcoded Partner="Amazon" for every ASIN row, reasoning
# that ASIN is an Amazon identifier. **That rule is empirically false in this
# system.** In the test environment, ASIN-type product IDs are tagged
# Amazon 1,559 / Other 536 / Novi 527, with 528 untagged — so ASIN does not
# imply Amazon.
#
# The defensible signal is the client's own cert sheet: of the 158 certificates
# in the ASIN master list, **154 carry cert-level Partner = Amazon**. So the tag
# is inherited from the certificate.
#
# The other 4 (23-0814 P&G, and Kohler 25-8535 / 26-0053 / 26-0101) have a BLANK
# cert-level Partner and between them carry **1,960 ASINs — 12% of the load**.
# They are left UNTAGGED and reported, because asserting "Amazon" there would
# invent a partner relationship the client never stated — the same reason
# fix_email() refuses to guess a missing TLD.
#
# ⚠️ Untagged is NOT neutral: under CR-171 Model A an untagged product ID is
# returned to EVERY partner caller (DEF-352). So the untagged 1,960 are
# fail-OPEN and need a client answer before go-live — they are listed in
# review-untagged-product-ids.csv rather than left to be discovered.

SHEET_ACTIVE = "Active List 7.31.26"
SHEET_EXPIRED2 = "Expired List 2 - 8.11.26"
SHEET_EXPIRED1 = "Expired List 1 - 4.9.26"

PREFIX_RE = re.compile(r"^\d+(?:\.\d+)*\s+")
ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")          # ProductIdService.cs:96
EXCEL_EPOCH = datetime.datetime(1899, 12, 30)

# Copied verbatim from CertificateImportService.cs:201 so this script rejects
# exactly what the importer rejects — no more, no less.
EMAIL_RE = re.compile(r"^[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}$")
ANGLE_EMAIL_RE = re.compile(r"<([^<>]+)>")


def fix_email(v, sheet, excel_row, number, review):
    """Normalise Company Email, and never let a bad one fail the whole certificate.

    Two real defects in the client's source, found by previewing certs-expired2:
      * Outlook display-name form 'Name <addr@x.com>' (8 rows) -> extract the
        address. Decision-free, so it is fixed here.
      * 'info@yidesigngroup' — no TLD (4 rows). NOT guessable ('.com' is a guess
        that silently invents client contact data), so the value is dropped and
        reported instead.

    A malformed optional contact field must not cost us the certificate row —
    but dropping it silently would be the very thing DEF-348 is about, so every
    drop lands in review-invalid-company-emails.csv.
    """
    s = clean(v)
    if not s:
        return ""
    m = ANGLE_EMAIL_RE.search(s)
    if m:
        inner = clean(m.group(1))
        if EMAIL_RE.match(inner):
            if inner.casefold() != s.casefold():
                review.append({"Sheet": sheet, "ExcelRow": excel_row, "CertificateNumber": number,
                               "Original": s, "Emitted": inner,
                               "Action": "extracted address from display-name form"})
            return inner
        s = inner
    if EMAIL_RE.match(s):
        return s
    review.append({"Sheet": sheet, "ExcelRow": excel_row, "CertificateNumber": number,
                   "Original": clean(v), "Emitted": "",
                   "Action": "DROPPED - fails the importer's email regex; not mechanically fixable"})
    return ""


# ------------------------------------------------------------------ helpers --
def clean(v) -> str:
    """Trim, collapse whitespace, normalise NBSP. Never returns None."""
    if v is None:
        return ""
    s = str(v).replace(" ", " ").strip()
    return re.sub(r"\s+", " ", s)


def strip_prefix(v) -> str:
    """'4.3 Cleaning Supplies' -> 'Cleaning Supplies'. The numeric prefix is a
    spreadsheet grouping aid; the client states it must not reach the DB."""
    return PREFIX_RE.sub("", clean(v)).strip()


def to_date(v) -> str:
    """-> 'yyyy-MM-dd'. Pre-formatted deliberately: UtcDateTimeConverter shifts
    Unspecified-Kind dates, so calendar fields must be strings by the time they
    reach the importer. Handles real datetimes, Excel serials, and text."""
    if v is None or clean(v) == "":
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        # Excel serial (the remediation script's EXP-DATES rule). Guard the range:
        # 1 = 1899-12-31, 60000 ~= 2064. Anything outside is not a date.
        if 1 <= float(v) <= 60000:
            return (EXCEL_EPOCH + datetime.timedelta(days=float(v))).strftime("%Y-%m-%d")
        return ""
    s = clean(v)
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return s  # leave it visible rather than silently blanking it


def header_map(ws):
    """header text -> [column indexes]. Duplicated headers keep every index —
    the Certificate Subcategories column repeats 4-6x per sheet."""
    idx = defaultdict(list)
    for row in ws.iter_rows(min_row=1, max_row=1):
        for i, cell in enumerate(row):
            if cell.value is not None and clean(cell.value):
                idx[clean(cell.value)].append(i)
    return idx


def find_cols(idx, wanted):
    """Tolerant source-side lookup: the sheets carry trailing spaces
    ('Product Type ') and per-sheet parenthetical suffixes."""
    w = wanted.casefold()
    for h, cols in idx.items():
        if h.casefold() == w:
            return cols
    for h, cols in idx.items():
        if h.casefold().startswith(w):
            return cols
    return []


def get(row, cols, i=0):
    if not cols or i >= len(cols):
        return ""
    ci = cols[i]
    return row[ci] if ci < len(row) else ""


# --------------------------------------------------------------- transform --
def transform_sheet(ws, sheet_name, has_new_fields, unmappable, pathway_fix, email_review):
    """Yield one dict per source row, plus the row's Excel line number."""
    idx = header_map(ws)

    c_num = find_cols(idx, "CERTIFICATE_NUMBER")
    c_co = find_cols(idx, "COMPANY")
    c_prod = find_cols(idx, "PRODUCT")
    c_type = find_cols(idx, "CERTIFICATE_TYPE")
    c_eff = find_cols(idx, "EFFECTIVE_DATE")
    c_exp = find_cols(idx, "EXPIRATION_DATE")
    c_status = find_cols(idx, "STATUS")

    if has_new_fields:
        c_cat = find_cols(idx, "Product Category")
        c_ptype = find_cols(idx, "Product Type")
        c_psub = find_cols(idx, "Product Subcategory")
        c_loc = find_cols(idx, "Location (no country)")
        c_country = find_cols(idx, "Country")
        c_partner = find_cols(idx, "Partner (Amazon)")
        c_results = find_cols(idx, "Certification Results")
        c_pathway = find_cols(idx, "Pathway (SEP Amazon only)")
        c_extag = find_cols(idx, "External Tags")
        c_intag = find_cols(idx, "Internal Tags")
        c_email = find_cols(idx, "Contact Email")
        c_url = find_cols(idx, "Contact URL")
        c_certsub = find_cols(idx, "Certificate Subcategories")
    else:
        # Expired List 1 has none of the new-schema fields. The client is explicit:
        # "New DB fields not listed in tab should be left blank we did not assign".
        # So Country stays BLANK here — we do not infer it out of LOCATION.
        c_cat = c_ptype = c_psub = c_country = c_partner = []
        c_results = c_pathway = c_extag = c_intag = c_email = c_url = c_certsub = []
        c_loc = find_cols(idx, "LOCATION")

    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(c is None or clean(c) == "" for c in row):
            continue

        number = clean(get(row, c_num))
        if not number:
            continue  # blank cert# — reported by the validation gate, not loadable

        # Certificate Subcategories repeats across several columns; join non-empties.
        certsubs = ""
        if c_certsub:
            vals = [clean(get(row, c_certsub, i)) for i in range(len(c_certsub))]
            certsubs = ", ".join(v for v in vals if v)

        subcat = strip_prefix(get(row, c_psub)) if c_psub else ""

        # DEF-348 — count what the parser will drop on the floor.
        for field, value in (("Subcategory", subcat), ("Cert Subcategories", certsubs)):
            if value:
                unmappable.append(
                    {
                        "Sheet": sheet_name,
                        "ExcelRow": excel_row,
                        "CertificateNumber": number,
                        "Field": field,
                        "Value": value,
                        "Why": "no parser mapping in CertificateImportService (DEF-348)",
                    }
                )

        pathway = clean(get(row, c_pathway))
        if pathway:
            pathway = pathway_fix.get(pathway.casefold(), pathway)

        yield excel_row, {
            "Certificate #": number,
            "Certification Type": clean(get(row, c_type)),
            "Cert Subcategories": certsubs,
            "Status": clean(get(row, c_status)),
            "Company": clean(get(row, c_co)),
            "Company Email": fix_email(get(row, c_email), sheet_name, excel_row, number, email_review),
            "Company URL": clean(get(row, c_url)),
            "Product Name": clean(get(row, c_prod)),
            "Product Category": strip_prefix(get(row, c_cat)) if c_cat else "",
            "Product Type": strip_prefix(get(row, c_ptype)) if c_ptype else "",
            "Subcategory": subcat,
            "Partner": clean(get(row, c_partner)),
            "Pathway": pathway,
            "Cert Results": clean(get(row, c_results)),
            "Effective Date": to_date(get(row, c_eff)),
            "Expiration Date": to_date(get(row, c_exp)),
            "Location": clean(get(row, c_loc)),
            "Country": clean(get(row, c_country)),
            "External Tag": clean(get(row, c_extag)),
            "Internal Tag": clean(get(row, c_intag)),
        }


def completeness(rec) -> int:
    return sum(1 for v in rec.values() if clean(v))


# A row missing any of these cannot become a certificate under ANY loader:
#   Company / Product Name / Certification Type -> the importer rejects them
#     ("X is required"), and Certificates.CompanyId / ProductId /
#     CertificationTypeId are all NOT NULL.
#   Effective / Expiration Date -> DEF-353. The importer used to substitute
#     today / today+1yr, which produced an Expired 2013 certificate with a 2027
#     expiry and two certs whose effective date fell AFTER their expiration.
#     Both loaders now refuse: the importer rejects in preview, the SQL loader
#     aborts in preflight.
REQUIRED_FIELDS = ["Company", "Product Name", "Certification Type",
                   "Effective Date", "Expiration Date"]


def withhold_incomplete(rows, sheet_name, quarantine):
    """Hold back rows that cannot load, and say why.

    Withholding HERE rather than letting the loader fail is the point: the load set
    stays clean and every excluded row is reported for the client to complete —
    the same treatment duplicate certificate numbers get. Discovered because the
    SQL loader's own assertion caught 5,390 inserted against 5,401 staged and
    rolled back; the 11 were blank Product Names."""
    keep = []
    for excel_row, rec in rows:
        missing = [f for f in REQUIRED_FIELDS if not rec[f].strip()]
        if missing:
            quarantine.append({
                "Sheet": sheet_name,
                "ExcelRow": excel_row,
                "CertificateNumber": rec["Certificate #"],
                "Company": rec["Company"],
                "ProductName": rec["Product Name"],
                "Status": rec["Status"],
                "MissingFields": ", ".join(missing),
            })
        else:
            keep.append((excel_row, rec))
    return keep


def apply_dup_policy(rows, policy, sheet_name, quarantine):
    """rows: list of (excel_row, rec). Returns the rows to emit."""
    by_num = defaultdict(list)
    for excel_row, rec in rows:
        by_num[rec["Certificate #"]].append((excel_row, rec))

    emit = []
    for number, group in by_num.items():
        if len(group) == 1:
            emit.append(group[0])
            continue

        if policy == "all":
            emit.extend(group)
            for excel_row, rec in group:
                quarantine.append(_q(sheet_name, excel_row, rec, "emitted anyway (--dup-policy all); WILL violate IX_Certificates_Number"))
        elif policy == "collapse":
            group_sorted = sorted(group, key=lambda t: (-completeness(t[1]), t[0]))
            keeper = group_sorted[0]
            emit.append(keeper)
            quarantine.append(_q(sheet_name, keeper[0], keeper[1], "KEPT (most complete of %d)" % len(group)))
            for excel_row, rec in group_sorted[1:]:
                quarantine.append(_q(sheet_name, excel_row, rec, "DROPPED by collapse — its product is not loaded"))
        else:  # quarantine
            for excel_row, rec in group:
                quarantine.append(_q(sheet_name, excel_row, rec, "HELD — duplicate cert# decision open (plan 11.2)"))

    emit.sort(key=lambda t: t[0])
    return emit


def _q(sheet_name, excel_row, rec, disposition):
    return {
        "Sheet": sheet_name,
        "ExcelRow": excel_row,
        "CertificateNumber": rec["Certificate #"],
        "Company": rec["Company"],
        "ProductName": rec["Product Name"],
        "CertificationType": rec["Certification Type"],
        "EffectiveDate": rec["Effective Date"],
        "ExpirationDate": rec["Expiration Date"],
        "Disposition": disposition,
    }


def write_csv(path, headers, rows):
    # utf-8-sig: the importer's own template is written with a BOM so Excel
    # decodes UTF-8 rather than Windows-1252 (CertificateImportController:468).
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=headers, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


# -------------------------------------------------------------- transform C --
def transform_asins(path, outdir, loadable_numbers, summary, cert_partner):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    idx = header_map(ws)

    c_type = find_cols(idx, "type")
    c_value = find_cols(idx, "value")
    c_code = find_cols(idx, "amazonCode")
    c_pname = find_cols(idx, "productName")
    c_cert = find_cols(idx, "certificateNumber")
    c_pathway = find_cols(idx, "pathway")

    seen = set()
    per_cert = defaultdict(list)
    invalid, orphans = [], []
    dupes = 0
    total = 0
    pathway_case = Counter()

    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(c is None or clean(c) == "" for c in row):
            continue
        total += 1

        cert = clean(get(row, c_cert))
        idtype = clean(get(row, c_type)).upper()
        value = clean(get(row, c_value)).upper()
        if c_pathway:
            pv = clean(get(row, c_pathway))
            if pv:
                pathway_case[pv] += 1

        if idtype == "ASIN" and not ASIN_RE.match(value):
            invalid.append({"ExcelRow": excel_row, "CertificateNumber": cert, "Type": idtype,
                            "Value": value, "Why": "fails ProductIdService ASIN pattern ^[A-Z0-9]{10}$"})
            continue

        # DB enforces per-certificate uniqueness (FixProductIdUniqueIndex_PerCertificate),
        # so 4,797 surplus (cert#, ASIN) pairs must collapse here rather than fail on insert.
        key = (cert, idtype, value)
        if key in seen:
            dupes += 1
            continue
        seen.add(key)

        if cert not in loadable_numbers:
            orphans.append({"ExcelRow": excel_row, "CertificateNumber": cert, "Type": idtype,
                            "Value": value,
                            "Why": "certificate not in the emitted load set (quarantined dup, or absent)"})
            continue

        per_cert[cert].append({
            "type": idtype,
            "value": value,
            # inherited from the certificate; blank when the cert names no partner
            "Partner": cert_partner.get(cert, ""),
            "productName": clean(get(row, c_pname)),
        })

    piddir = os.path.join(outdir, "productids")
    os.makedirs(piddir, exist_ok=True)
    for cert, items in sorted(per_cert.items()):
        safe = re.sub(r"[^A-Za-z0-9._-]", "_", cert)
        write_csv(os.path.join(piddir, f"{safe}.csv"), PRODUCTID_HEADERS, items)

    combined = []
    for cert, items in sorted(per_cert.items()):
        for it in items:
            combined.append({"certificateNumber": cert, **it})
    write_csv(os.path.join(outdir, "productids-all.csv"),
              ["certificateNumber"] + PRODUCTID_HEADERS, combined)

    if invalid:
        write_csv(os.path.join(outdir, "quarantine-invalid-asins.csv"),
                  ["ExcelRow", "CertificateNumber", "Type", "Value", "Why"], invalid)
    if orphans:
        write_csv(os.path.join(outdir, "quarantine-orphan-asins.csv"),
                  ["ExcelRow", "CertificateNumber", "Type", "Value", "Why"], orphans)

    print(f"  source rows            {total}")
    print(f"  invalid ASIN values    {len(invalid)}  -> quarantine-invalid-asins.csv")
    print(f"  surplus duplicate pairs{dupes:>5}  (collapsed; DB enforces per-cert uniqueness)")
    print(f"  orphaned (cert not loadable) {len(orphans)}  -> quarantine-orphan-asins.csv")
    untagged = [r for r in combined if not r["Partner"]]
    if untagged:
        write_csv(os.path.join(outdir, "review-untagged-product-ids.csv"),
                  ["certificateNumber"] + PRODUCTID_HEADERS, untagged)
        certs_u = sorted({r["certificateNumber"] for r in untagged})
        print(f"  !! UNTAGGED (no cert-level Partner): {len(untagged)} IDs across {len(certs_u)} certs "
              f"-> review-untagged-product-ids.csv")
        print(f"     certs: {', '.join(certs_u)}")
        print(f"     CR-171 Model A: untagged = visible to EVERY partner (DEF-352). Needs a client answer.")
    print(f"  emitted                {len(combined)} product IDs across {len(per_cert)} certificates")
    if len(pathway_case) > 1:
        print(f"  NOTE pathway case variants in source: {dict(pathway_case)}")

    summary.append({"Stage": "C-asins", "Sheet": ws.title, "SourceRows": total,
                    "Emitted": len(combined), "Quarantined": len(invalid) + len(orphans),
                    "Note": f"{len(per_cert)} certs; {dupes} surplus dup pairs collapsed"})


# -------------------------------------------------------------------- main --
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("certs", nargs="?", default=DEFAULT_CERTS)
    ap.add_argument("asins", nargs="?", default=DEFAULT_ASINS)
    ap.add_argument("--dup-policy", choices=["quarantine", "collapse", "all"], default="quarantine")
    args = ap.parse_args()

    for p in (args.certs, args.asins):
        if not os.path.exists(p):
            sys.exit(f"Source not found: {p}")

    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    outdir = os.path.join(HERE, "_transform_output", stamp)
    os.makedirs(outdir, exist_ok=True)

    print(f"dup-policy = {args.dup_policy}")
    print(f"output     = {outdir}\n")

    # Canonical pathway casing, so "Renewable Electricity Usage" matches the seeded row.
    pathway_fix = {
        "renewable electricity usage": "Renewable electricity usage",
        "energy usage reduction": "Energy usage reduction",
        "onsite fossil fuel usage reduction": "Onsite fossil fuel usage reduction",
    }

    wb = openpyxl.load_workbook(args.certs, read_only=True, data_only=True)

    unmappable = []
    email_review = []
    date_quarantine = []
    quarantine = []
    summary = []
    loadable_numbers = set()
    cert_partner = {}

    sheets = [
        (SHEET_ACTIVE, "certs-active.csv", True),
        (SHEET_EXPIRED2, "certs-expired2.csv", True),
        (SHEET_EXPIRED1, "certs-expired1.csv", False),
    ]

    for sheet_name, outfile, has_new in sheets:
        if sheet_name not in wb.sheetnames:
            print(f"!! sheet missing: {sheet_name}")
            continue
        ws = wb[sheet_name]
        rows = list(transform_sheet(ws, sheet_name, has_new, unmappable, pathway_fix, email_review))
        rows = withhold_incomplete(rows, sheet_name, date_quarantine)
        emit = apply_dup_policy(rows, args.dup_policy, sheet_name, quarantine)

        write_csv(os.path.join(outdir, outfile), CERT_HEADERS, [r for _, r in emit])
        for _, rec in emit:
            loadable_numbers.add(rec["Certificate #"])
            cert_partner[rec["Certificate #"]] = (rec["Partner"] or "").strip()

        held = len(rows) - len(emit)
        print(f"== {sheet_name}")
        print(f"   source rows {len(rows):5}   emitted {len(emit):5}   held back {held:5}   -> {outfile}")
        summary.append({"Stage": "A/B-certs", "Sheet": sheet_name, "SourceRows": len(rows),
                        "Emitted": len(emit), "Quarantined": held,
                        "Note": f"dup-policy={args.dup_policy}"})

    print("\n== ASIN master list (Transform C)")
    transform_asins(args.asins, outdir, loadable_numbers, summary, cert_partner)

    if quarantine:
        write_csv(os.path.join(outdir, "quarantine-duplicate-cert-numbers.csv"),
                  ["Sheet", "ExcelRow", "CertificateNumber", "Company", "ProductName",
                   "CertificationType", "EffectiveDate", "ExpirationDate", "Disposition"],
                  quarantine)
    if unmappable:
        write_csv(os.path.join(outdir, "unmappable-fields.csv"),
                  ["Sheet", "ExcelRow", "CertificateNumber", "Field", "Value", "Why"], unmappable)
    if date_quarantine:
        write_csv(os.path.join(outdir, "quarantine-incomplete-rows.csv"),
                  ["Sheet", "ExcelRow", "CertificateNumber", "Company", "ProductName",
                   "Status", "MissingFields"], date_quarantine)
    if email_review:
        write_csv(os.path.join(outdir, "review-invalid-company-emails.csv"),
                  ["Sheet", "ExcelRow", "CertificateNumber", "Original", "Emitted", "Action"], email_review)

    write_csv(os.path.join(outdir, "transform-summary.csv"),
              ["Stage", "Sheet", "SourceRows", "Emitted", "Quarantined", "Note"], summary)

    print("\n---------------------------------------------------------------")
    print(f"cert rows quarantined (dup cert#) : {len(quarantine)}")
    print(f"values with NO import path (DEF-348): {len(unmappable)}")
    print(f"withheld - incomplete rows          : {len(date_quarantine)}")
    fixed = sum(1 for r in email_review if r["Emitted"])
    print(f"company emails repaired / dropped  : {fixed} / {len(email_review) - fixed}")
    print(f"loadable certificate numbers       : {len(loadable_numbers)}")
    print(f"\nOutput: {outdir}")


if __name__ == "__main__":
    main()
