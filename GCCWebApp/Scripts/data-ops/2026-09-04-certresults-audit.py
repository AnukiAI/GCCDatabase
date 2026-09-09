# -*- coding: utf-8 -*-
"""
CERTIFICATION RESULTS AUDIT  --  source workbooks vs the database        CR-045
==============================================================================
Answers one question per certificate: does the value stored in the database match
what the CLIENT SEES in their own spreadsheet cell?

THE WHOLE POINT -- displayed value, not raw cell value.
  Every populated "Certification Results" cell in these workbooks carries the Excel
  number format '0%'.  Such a cell HOLDS 0.23 and DISPLAYS "23%".  Reading the raw
  value is what put bare decimals in the database in the first place; CR-042 then
  rewrote them to explicit "XX%".  So the audit compares against the DISPLAYED
  value, and a row whose DB value equals the RAW decimal is reported as a MISS.

  Text cells are the exception and are handled as such: a number format does not
  apply to a string, so "99% Post-Consumer" displays verbatim.

Reads nothing but SELECTs.  Writes one CSV.  Changes no data.

    python 2026-09-04-certresults-audit.py test
    python 2026-09-04-certresults-audit.py dev

There is no default target on purpose -- two scripts in this folder once silently
defaulted to dev and produced findings about the wrong environment.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os
import sys
import warnings
from collections import defaultdict

warnings.filterwarnings("ignore")
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl
import pyodbc

HERE = os.path.dirname(os.path.abspath(__file__))
OUTBASE = os.path.join(HERE, "_certresults_audit")

MIGDIR = (r"C:\Users\swest\OneDrive - PointSolve Technology Inc\LocalWest"
          r"\Projects\GCC Database Solution\DataMigrationTestDev")
SOURCES = [
    os.path.join(MIGDIR, "GCC Database Migration_Certs_8.12.26V2.xlsx"),
    os.path.join(MIGDIR, "new", "GCC Database Migration_Certs_9.3.26.xlsx"),
]

TARGETS = {
    "dev": ("Driver={ODBC Driver 17 for SQL Server};Server=ANUKI01;Database=GCCDatabase_Dev;"
            "Uid=GCC_DevUser;Pwd=DevP@ss2024!SQL;TrustServerCertificate=yes;Encrypt=no;"),
    "test": ("Driver={ODBC Driver 17 for SQL Server};Server=tcp:68.71.129.120,1533;"
             "Database=gccwebap_certificates;Uid=GCC_adminxyz123;Pwd=Q3Ez4e!gE_47KNi;"
             "Encrypt=yes;TrustServerCertificate=yes;"),
}

if len(sys.argv) < 2 or sys.argv[1] not in TARGETS:
    sys.exit("usage: 2026-09-04-certresults-audit.py {dev|test}")
target = sys.argv[1]

stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
outdir = os.path.join(OUTBASE, stamp)
os.makedirs(outdir, exist_ok=True)
OUTCSV = os.path.join(outdir, "audit.csv")

print("### CERTIFICATION RESULTS AUDIT   target=%s   %s" % (target.upper(), stamp))


# ---------------------------------------------------------------- Excel side --
def decimals_in_format(fmt):
    """Decimal places a numeric Excel number format renders. '0%' -> 0, '0.00%' -> 2."""
    body = fmt.split(";")[0]
    if "." not in body:
        return 0
    frac = body.split(".", 1)[1]
    return sum(1 for ch in frac if ch in "0#")


def displayed_value(value, fmt):
    """What Excel actually puts on screen. Returns (displayed, hides_precision, kind)."""
    if isinstance(value, bool):
        return str(value), False, "bool"
    if isinstance(value, str):
        # A number format does not apply to text -- the cell shows the string itself.
        return value.strip(), False, "text"
    if isinstance(value, (int, float)):
        if "%" in fmt:
            dec = decimals_in_format(fmt)
            pct = value * 100
            shown = "%.*f%%" % (dec, pct)
            hidden = abs(pct - round(pct, dec)) > 1e-9
            return shown, hidden, "percent"
        return ("%g" % value), False, "number"
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d"), False, "date"
    return str(value), False, type(value).__name__


source_rows = []          # one entry per populated source cell
source_numbers = set()    # every cert number seen in any source sheet
source_blank = {}         # number -> (file, sheet) where the results cell was blank

for path in SOURCES:
    fname = os.path.basename(path)
    wb = openpyxl.load_workbook(path, data_only=True)   # not read_only: need number_format
    for sn in wb.sheetnames:
        ws = wb[sn]
        hdr = [(c.value if c.value is not None else "") for c in ws[1]]
        ci = [i for i, h in enumerate(hdr) if "certification result" in str(h).lower()]
        ni = [i for i, h in enumerate(hdr) if "certificate_number" in str(h).lower()]
        if not ni:
            continue
        # Register every cert number the workbooks mention, even on sheets that carry
        # no results column -- otherwise the db-only direction reports certificates as
        # "not in either source workbook" when they are simply listed on another tab.
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=ni[0] + 1, max_col=ni[0] + 1):
            v = row[0].value
            if v is not None and str(v).strip():
                source_numbers.add(str(v).strip().upper())
        if not ci:
            continue
        ci, ni = ci[0], ni[0]
        print("  reading %s :: '%s'  (results col %d, %d data rows)"
              % (fname, sn, ci + 1, ws.max_row - 1))
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            num = row[ni].value
            if num is None or str(num).strip() == "":
                continue
            num = str(num).strip().upper()
            source_numbers.add(num)
            cell = row[ci]
            if cell.value is None or str(cell.value).strip() == "":
                source_blank.setdefault(num, (fname, sn))
                continue
            shown, hidden, kind = displayed_value(cell.value, cell.number_format)
            source_rows.append({
                "file": fname, "sheet": sn, "number": num,
                "raw": str(cell.value).strip(), "type": kind,
                "fmt": cell.number_format, "shown": shown, "hidden": hidden,
            })
    wb.close()

print("  -> %d populated source cells across %d cert numbers"
      % (len(source_rows), len(source_numbers)))


# ------------------------------------------------------------------- DB side --
cn = pyodbc.connect(TARGETS[target])
cur = cn.cursor()
cur.execute("SELECT Number, Id, CertificationResults, CertificationResultsNumeric FROM dbo.Certificates")
db = defaultdict(list)
for number, cid, res, numeric in cur.fetchall():
    if numeric is None:
        nstr = ""
    else:
        nstr = format(numeric.normalize(), "f") if hasattr(numeric, "normalize") else str(numeric)
    db[str(number).strip().upper()].append({"id": cid, "res": (res or "").strip(), "num": nstr})
cur.close()
cn.close()

db_total = sum(len(v) for v in db.values())
print("  -> %d certificate rows in %s across %d distinct numbers"
      % (db_total, target.upper(), len(db)))


# ------------------------------------------------------------------ compare --
def norm(s):
    return s.strip().rstrip("%").strip() if s else ""


out = []
tally = defaultdict(int)

for s in source_rows:
    rows = db.get(s["number"], [])
    stored = [r["res"] for r in rows]
    populated = [v for v in stored if v]
    distinct = sorted(set(stored))

    if not rows:
        verdict, note = "NOT_IN_DB", "certificate number absent from this environment"
    elif not populated:
        verdict, note = "DB_BLANK", "source has a value, database field is empty"
    elif len(set(populated)) > 1:
        verdict = "DB_ROWS_DISAGREE"
        note = "%d rows, %d different values" % (len(rows), len(set(populated)))
    elif populated[0] == s["shown"]:
        verdict, note = "OK", ""
    elif populated[0] == s["raw"]:
        verdict, note = "MISMATCH_STORED_RAW", "DB holds the raw cell value, not what the client sees"
    elif norm(populated[0]) == norm(s["shown"]):
        verdict, note = "OK_FORMATTING", "same number, different punctuation or percent sign"
    else:
        verdict, note = "MISMATCH_OTHER", "DB value matches neither the displayed nor the raw cell value"

    if s["hidden"] and verdict.startswith("OK"):
        note = (note + "; " if note else "") + "Excel format hides precision: cell holds " + s["raw"]

    tally[verdict] += 1
    out.append({
        "direction": "source->db",
        "source_file": s["file"], "source_sheet": s["sheet"],
        "certificate_number": s["number"],
        "source_raw_cell_value": s["raw"], "source_cell_type": s["type"],
        "source_number_format": s["fmt"], "source_displayed_to_client": s["shown"],
        "db_row_count": len(rows),
        "db_stored_value": " | ".join(distinct),
        "db_sort_numeric": " | ".join(sorted(set(r["num"] for r in rows))),
        "verdict": verdict, "note": note,
    })

# The other direction: DB values whose cert number appears in NO source sheet.
for number, rows in sorted(db.items()):
    if number in source_numbers:
        continue
    populated = [r["res"] for r in rows if r["res"]]
    if not populated:
        continue
    tally["DB_ONLY"] += 1
    out.append({
        "direction": "db-only",
        "source_file": "", "source_sheet": "", "certificate_number": number,
        "source_raw_cell_value": "", "source_cell_type": "",
        "source_number_format": "", "source_displayed_to_client": "",
        "db_row_count": len(rows),
        "db_stored_value": " | ".join(sorted(set(populated))),
        "db_sort_numeric": " | ".join(sorted(set(r["num"] for r in rows))),
        "verdict": "DB_ONLY",
        "note": "value in the database, cert number not in either source workbook",
    })

# A database value with no source cell to check it against -- the certificate IS named
# in the workbooks, but only on a sheet that carries no Certification Results column.
# Reported rather than dropped: unverifiable is a finding, not a pass.
populated_numbers = set(r["number"] for r in source_rows)
for number, rows in sorted(db.items()):
    if number not in source_numbers or number in populated_numbers or number in source_blank:
        continue
    populated = [r["res"] for r in rows if r["res"]]
    if not populated:
        continue
    tally["NO_SOURCE_CELL"] += 1
    out.append({
        "direction": "db-only",
        "source_file": "", "source_sheet": "(sheet has no results column)",
        "certificate_number": number,
        "source_raw_cell_value": "", "source_cell_type": "",
        "source_number_format": "", "source_displayed_to_client": "",
        "db_row_count": len(rows),
        "db_stored_value": " | ".join(sorted(set(populated))),
        "db_sort_numeric": " | ".join(sorted(set(r["num"] for r in rows))),
        "verdict": "NO_SOURCE_CELL",
        "note": "database holds a value; the workbooks list this cert only on a sheet without a results column",
    })

# Source cell blank but the database holds a value.
for number, meta in sorted(source_blank.items()):
    if number in populated_numbers:
        continue   # populated on some other sheet in the workbook set
    rows = db.get(number, [])
    populated = [r["res"] for r in rows if r["res"]]
    if not populated:
        continue
    tally["SOURCE_BLANK_DB_HAS"] += 1
    out.append({
        "direction": "db-only",
        "source_file": meta[0], "source_sheet": meta[1], "certificate_number": number,
        "source_raw_cell_value": "", "source_cell_type": "(blank)",
        "source_number_format": "", "source_displayed_to_client": "",
        "db_row_count": len(rows),
        "db_stored_value": " | ".join(sorted(set(populated))),
        "db_sort_numeric": " | ".join(sorted(set(r["num"] for r in rows))),
        "verdict": "SOURCE_BLANK_DB_HAS",
        "note": "source cell is empty, database has a value",
    })

FIELDS = ["direction", "source_file", "source_sheet", "certificate_number",
          "source_raw_cell_value", "source_cell_type", "source_number_format",
          "source_displayed_to_client", "db_row_count", "db_stored_value",
          "db_sort_numeric", "verdict", "note"]

ORDER = {"MISMATCH_STORED_RAW": 0, "MISMATCH_OTHER": 1, "DB_ROWS_DISAGREE": 2,
         "DB_BLANK": 3, "NOT_IN_DB": 4, "SOURCE_BLANK_DB_HAS": 5,
         "DB_ONLY": 6, "NO_SOURCE_CELL": 7, "OK_FORMATTING": 8, "OK": 9}
out.sort(key=lambda r: (ORDER.get(r["verdict"], 9), r["certificate_number"]))

with io.open(OUTCSV, "w", encoding="utf-8-sig", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=FIELDS)
    w.writeheader()
    w.writerows(out)

print("\n### RESULT  (%d rows written)" % len(out))
for v in sorted(tally, key=lambda k: ORDER.get(k, 9)):
    print("    %-22s %5d" % (v, tally[v]))
print("\n### audit.csv -> %s" % OUTCSV)

bad = sum(tally[v] for v in ("MISMATCH_STORED_RAW", "MISMATCH_OTHER", "DB_ROWS_DISAGREE"))
print("### %d row(s) where the database disagrees with what the client sees." % bad)
