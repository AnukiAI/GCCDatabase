#!/usr/bin/env python3
"""
Client email (markdown)  ->  formatted Word document.
=====================================================

Reusable converter for the client-email drafts in docs/tracking/. Replaces hand-written
per-email builders: write the email as markdown, run this, send the .docx.

WHAT IT REFUSES TO INCLUDE
--------------------------
Our email drafts carry internal material -- "read before sending" warnings, file:line evidence,
notes addressed to Shawn. A Word file gets forwarded, so the converter treats the client-facing
body as a strict subset and everything outside it is dropped:

  * Everything BEFORE the `**Subject:**` line   (title, Status/Attach metadata, internal callouts)
  * Everything FROM a `# Not part of the email...` heading onward
  * Everything FROM a `---` / `---` double divider onward

If a draft has no `**Subject:**` line the converter STOPS rather than guessing where the body
begins -- failing loudly beats silently shipping an internal note.

SUPPORTED MARKDOWN
------------------
  ## Heading          section heading (rule beneath)
  ### Heading         sub-heading
  **bold**            inline bold
  `code`              inline, rendered as plain text (no scary monospace for a client)
  | a | b |           tables, first row treated as header
  - item              bullets
  > text              callout box
  ---                 ignored
  <!-- pagebreak -->  start a new page here (author-controlled; invisible in markdown)

Usage
  python md-to-client-letter.py <input.md> [output.docx] [--break-before 2,4]

  --break-before N,M   start a new page before the "## N." and "## M." sections. Layout choice
                       kept out of the draft, so the markdown stays purely content.

  python md-to-client-letter.py docs/tracking/2026-08-18-client-email-migration-complete.md \
      --break-before 2,4
"""
import os
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches
except ImportError:
    sys.exit("python-docx required:  pip install python-docx")

PINE  = RGBColor(0x1F, 0x4D, 0x38)
INK   = RGBColor(0x1A, 0x1A, 0x1A)
SLATE = RGBColor(0x4A, 0x55, 0x50)
BODY  = "Calibri"
HEAD  = "Georgia"
FIRM  = "ANUKI CONSULTING"
SUBFIRM = "GreenCircle Certified — Database Project"


# ------------------------------------------------------------------ docx helpers
def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def rule(paragraph, color="1F4D38", size=8):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pPr.append(borders)


def left_bar(paragraph, color="1F4D38"):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    pPr.append(borders)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text, size=10.5, color=INK, base_bold=False):
    """Render **bold**, *italic* and `code` inside a paragraph."""
    for part in INLINE.split(text):
        if not part:
            continue
        bold, italic = base_bold, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part = part[1:-1]                    # keep the text, drop the backticks
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        r = paragraph.add_run(part)
        r.font.name, r.font.size = BODY, Pt(size)
        r.font.bold, r.font.italic = bold, italic
        r.font.color.rgb = color


# ------------------------------------------------------------------ markdown slicing
def client_body(md_text, path):
    lines = md_text.splitlines()

    start = None
    subject = ""
    for i, ln in enumerate(lines):
        m = re.match(r"^\*\*Subject:\*\*\s*(.+)$", ln.strip())
        if m:
            subject = m.group(1).strip()
            start = i + 1
            break
    if start is None:
        sys.exit(f"ERROR: no '**Subject:**' line in {os.path.basename(path)}.\n"
                 "       Refusing to guess where the client-facing body begins — an internal note "
                 "could end up in the Word file.")

    end = len(lines)
    for i in range(start, len(lines)):
        if re.match(r"^#\s+Not part of the email", lines[i].strip(), re.I):
            end = i
            break
        # a '---' immediately followed by another '---' marks the internal appendix
        if lines[i].strip() == "---":
            nxt = next((l for l in lines[i + 1:] if l.strip()), "")
            if nxt.strip() == "---":
                end = i
                break

    return subject, lines[start:end]


# ------------------------------------------------------------------ render
def build(md_path, out_path, break_before=()):
    with open(md_path, encoding="utf-8") as f:
        subject, lines = client_body(f.read(), md_path)
    break_before = {str(b).strip() for b in break_before if str(b).strip()}

    doc = Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Inches(0.9), Inches(0.9)
        s.left_margin, s.right_margin = Inches(1.0), Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    def para(text="", size=10.5, bold=False, italic=False, color=INK, after=8, font=BODY):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        if text:
            r = p.add_run(text)
            r.font.name, r.font.size = font, Pt(size)
            r.font.bold, r.font.italic = bold, italic
            r.font.color.rgb = color
        return p

    # letterhead
    para(FIRM, size=9, bold=True, color=PINE, after=0)
    para(SUBFIRM, size=9, color=SLATE, after=14)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(subject)
    r.font.name, r.font.size, r.font.bold = HEAD, Pt(19), True
    r.font.color.rgb = PINE
    rule(p, size=12)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        # ---- explicit page break ----
        # Author-controlled rather than automatic: breaking before every section would
        # give a one-section-per-page document, which reads worse than a flowing letter.
        if re.match(r"^<!--\s*page-?break\s*-->$", line, re.I):
            doc.add_page_break()
            i += 1
            continue

        if not line or line == "---":
            i += 1
            continue

        # ---- fenced code block ----
        # There was no branch for ``` blocks, so they fell through to the inline-span parser: the
        # fences became LITERAL backticks and every line collapsed into one paragraph. That shipped
        # in a partner-facing letter. Render the block as monospace lines instead.
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1                                    # consume the closing fence
            while block and not block[0].strip():
                block.pop(0)
            while block and not block[-1].strip():
                block.pop()
            for bl in block:
                cp = doc.add_paragraph()
                cp.paragraph_format.space_after = Pt(0)
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.left_indent = Pt(18)
                # A blank line inside a block still needs a run, or Word collapses the spacing.
                r = cp.add_run(bl if bl.strip() else " ")
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # ---- table ----
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ci, h in enumerate(header):
                cell = t.rows[0].cells[ci]
                shade(cell, "1F4D38")
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(2)
                rr = cp.add_run(re.sub(r"\*\*|`", "", h))
                rr.font.name, rr.font.size, rr.font.bold = BODY, Pt(9.5), True
                rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for ri, row in enumerate(rows):
                cells = t.add_row().cells
                if ri % 2 == 1:
                    for c in cells:
                        shade(c, "F2F5F3")
                for ci, val in enumerate(row[:len(header)]):
                    cp = cells[ci].paragraphs[0]
                    cp.paragraph_format.space_after = Pt(2)
                    add_inline(cp, val, size=9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # ---- headings ----
        if line.startswith("### "):
            p = para(line[4:].strip(), size=11, bold=True, color=PINE, after=3)
            p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if line.startswith("## "):
            heading = line[3:].strip()
            num = re.match(r"^(\d+)[.)]", heading)
            if num and num.group(1) in break_before:
                doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading)
            r.font.name, r.font.size, r.font.bold = HEAD, Pt(14), True
            r.font.color.rgb = PINE
            rule(p)
            i += 1
            continue
        if line.startswith("# "):
            i += 1                      # the draft's own H1 is an internal title
            continue

        # ---- blockquote -> callout ----
        if line.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.left_indent = Inches(0.12)
            add_inline(p, " ".join(x for x in block if x), size=10)
            left_bar(p)
            continue

        # ---- bullets ----
        if re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, re.sub(r"^[-*]\s+", "", line))
            i += 1
            continue

        # ---- numbered list ----
        if re.match(r"^\d+[.)]\s+", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, re.sub(r"^\d+[.)]\s+", "", line))
            i += 1
            continue

        # ---- paragraph (join wrapped lines) ----
        buf = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", ">", "|", "---"))
                    or re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+[.)]\s+", nxt)):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        add_inline(p, " ".join(buf))

    # footer rule
    p = para("Anuki Consulting", size=9, color=SLATE, after=0)
    p.paragraph_format.space_before = Pt(18)
    rule(p, color="C9D4CE", size=6)

    try:
        doc.save(out_path)
    except PermissionError:
        alt = out_path.replace(".docx", "-new.docx")
        doc.save(alt)
        print(f"!! {out_path} is open in Word; wrote {alt} instead.")
        out_path = alt
    return out_path, subject


def main():
    argv = sys.argv[1:]
    breaks = []
    if "--break-before" in argv:
        k = argv.index("--break-before")
        if k + 1 >= len(argv):
            sys.exit("--break-before needs a value, e.g. --break-before 2,4")
        breaks = [x for x in re.split(r"[,\s]+", argv[k + 1]) if x]
        del argv[k:k + 2]

    if not argv:
        sys.exit(__doc__)
    md_path = os.path.abspath(argv[0])
    if not os.path.exists(md_path):
        sys.exit(f"not found: {md_path}")
    if len(argv) > 1:
        out = os.path.abspath(argv[1])
    else:
        base = os.path.basename(md_path)
        base = base.replace("client-email-", "client-letter-").replace(".md", ".docx")
        out = os.path.join(os.path.dirname(md_path), base)
    written, subject = build(md_path, out, breaks)
    print(f"subject: {subject}")
    if breaks:
        print(f"breaks : before section(s) {', '.join(breaks)}")
    print(f"wrote  : {written}")


if __name__ == "__main__":
    main()
