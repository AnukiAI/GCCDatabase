"""
Build the client-facing Word version of "Key System Backend Functions".

Source of truth for the wording:
    docs/tracking/2026-08-14-key-system-backend-functions.md

Content is authored here rather than parsed from the markdown so the Word output can carry
real styling (shaded table headers, callout blocks, page numbering). If the markdown changes,
change it here too.

Usage:  python build_backend_functions_docx.py <output.docx>
"""
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- palette
GREEN_DARK = RGBColor(0x1B, 0x5E, 0x3F)   # headings
GREEN_MID = RGBColor(0x2E, 0x7D, 0x54)    # sub-headings
INK = RGBColor(0x21, 0x21, 0x21)          # body
MUTED = RGBColor(0x5F, 0x6B, 0x66)        # metadata / captions
HDR_FILL = "1B5E3F"                       # table header shading
ALT_FILL = "F2F7F4"                       # zebra rows
NOTE_FILL = "FFF8E1"                      # callout
NOTE_RULE = "E0A800"


# ---------------------------------------------------------------- helpers
def shade(element, hex_fill):
    """Apply background shading to a cell or paragraph."""
    pr = element._tc.get_or_add_tcPr() if hasattr(element, "_tc") else element._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def cell_borders(cell, color="D5E0DA", size=4, edges=("top", "bottom", "left", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def no_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "left", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def left_accent(cell, color, size=18):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:color"), color)
    borders.append(el)
    tcPr.append(borders)


def cell_margins(cell, top=80, bottom=80, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def hrule(doc, color="D5E0DA", size=6, space_before=6, space_after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), color)
    borders.append(el)
    pPr.append(borders)
    return p


def page_number_footer(section, left_text):
    p = section.footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tab_stops = p.paragraph_format.tab_stops
    usable = section.page_width - section.left_margin - section.right_margin
    tab_stops.add_tab_stop(usable)

    r = p.add_run(left_text + "\t")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    r = p.add_run("Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    for instr in ("PAGE", "NUMPAGES"):
        run = p.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        itxt = OxmlElement("w:instrText")
        itxt.set(qn("xml:space"), "preserve")
        itxt.text = f" {instr} "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(itxt)
        run._r.append(end)
        if instr == "PAGE":
            sep = p.add_run(" of ")
            sep.font.size = Pt(8)
            sep.font.color.rgb = MUTED


def style_run(run, size=10.5, bold=False, color=INK, italic=False, mono=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    if mono:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def rich(paragraph, segments, size=10.5):
    """segments: list of (text, {bold, italic, mono, color})"""
    for text, opts in segments:
        r = paragraph.add_run(text)
        style_run(
            r,
            size=size,
            bold=opts.get("bold", False),
            italic=opts.get("italic", False),
            color=opts.get("color", INK),
            mono=opts.get("mono", False),
        )
    return paragraph


def body(doc, text=None, segments=None, size=10.5, space_after=8, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if segments:
        rich(p, segments, size=size)
    else:
        style_run(p.add_run(text), size=size)
    return p


def bullet(doc, segments, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    rich(p, segments, size=size)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    style_run(p.add_run(text), size=15, bold=True, color=GREEN_DARK)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    style_run(p.add_run(text), size=11.5, bold=True, color=GREEN_MID)
    return p


def make_table(doc, headers, rows, widths=None, zebra=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False

    hdr = t.rows[0]
    repeat_header(hdr)
    for i, htext in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        shade(c, HDR_FILL)
        cell_borders(c, color=HDR_FILL)
        cell_margins(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(htext), size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            c.text = ""
            cell_borders(c)
            cell_margins(c)
            if zebra and ri % 2 == 1:
                shade(c, ALT_FILL)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            segs = val if isinstance(val, list) else [(val, {})]
            rich(p, segs, size=10)

    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Inches(w)
    return t


def callout(doc, segments, fill=NOTE_FILL, rule=NOTE_RULE):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    c.text = ""
    shade(c, fill)
    left_accent(c, rule)
    cell_margins(c, top=130, bottom=130, left=170, right=140)
    c.width = Inches(6.5)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    rich(p, segments, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def code_block(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    c.text = ""
    shade(c, "F4F6F5")
    cell_borders(c, color="DDE4E0")
    cell_margins(c, top=120, bottom=120, left=150, right=140)
    c.width = Inches(6.5)
    first = True
    for line in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        style_run(p.add_run(line), size=9.5, mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------------------------------------------------------------- document
def build(out_path):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    s = doc.sections[0]
    s.top_margin = Inches(0.85)
    s.bottom_margin = Inches(0.85)
    s.left_margin = Inches(0.95)
    s.right_margin = Inches(0.95)
    page_number_footer(s, "GCC Database Solution  ·  Key System Backend Functions  ·  14 August 2026")

    # ---- title block
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run("GCC Database Solution"), size=10, bold=True, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    style_run(p.add_run("Key System Backend Functions"), size=24, bold=True, color=GREEN_DARK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    style_run(
        p.add_run("Reference for analyst training and user documentation"),
        size=11, color=MUTED, italic=True,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    rich(p, [
        ("Version 1.0", {"bold": True, "color": MUTED}),
        ("   ·   Updated 14 August 2026   ·   Verified against the application; all previously "
         "open items answered", {"color": MUTED}),
    ], size=9)

    hrule(doc, color=HDR_FILL, size=12, space_before=10, space_after=4)

    # ================================================ Certificate Number
    h1(doc, "Certificate Number")

    body(doc, segments=[
        ("Format:  ", {"bold": True}),
        ("YY-####", {"mono": True, "bold": True}),
        (" — two-digit year, hyphen, four-digit sequence. Example: ", {}),
        ("26-0001", {"mono": True}),
        (".", {}),
    ])

    bullet(doc, [
        ("Auto-assigned. ", {"bold": True}),
        ("The system allocates the next available number in sequence; analysts do not type it.", {}),
    ])
    bullet(doc, [
        ("Year rollover. ", {"bold": True}),
        ("The sequence restarts at ", {}), ("0001", {"mono": True}),
        (" each calendar year and the prefix advances. The first certificate created in 2027 is ", {}),
        ("27-0001", {"mono": True}),
        (". Nothing carries over from the prior year.", {}),
    ])
    bullet(doc, [
        ("Beyond 9,999. ", {"bold": True}),
        ("The number widens to five digits (e.g. ", {}), ("26-10000", {"mono": True}),
        (") only after all 9,999 numbers for that year have genuinely been used. ", {"bold": True}),
        ("Isolated high-numbered records — for example from an import or a legacy load — do not "
         "push new certificates into five digits.", {}),
    ])
    bullet(doc, [
        ("Assigned at save, not at form open. ", {"bold": True}),
        ("The number is allocated at the moment the record is written. If two analysts open the "
         "Create form at the same time, the first to save keeps the number they were shown; the "
         "second is silently assigned the next available number. ", {}),
        ("The same certificate number is never issued twice.", {"bold": True}),
    ])
    bullet(doc, [
        ("Field capacity: ", {"bold": True}),
        ("50 characters.", {}),
    ])

    body(doc, "", space_after=2)
    callout(doc, [
        ("Migrated records exception.  ", {"bold": True}),
        ("A limited set of records carried over from the prior system share a certificate number "
         "with another record, because those numbers were issued to more than one certificate "
         "before this system existed and cannot be changed after issue. These records are "
         "individually flagged. Newly created certificates can never take an existing number.", {}),
    ])

    # ================================================ Product Identifiers
    h1(doc, "Product Identifiers")

    body(doc, segments=[
        ("Where changes are saved differs by page — this is the point most often misunderstood.",
         {"bold": True}),
    ])

    make_table(
        doc,
        ["Page", "Behaviour"],
        [
            [[("Certificate → Edit", {"bold": True})],
             [("Product ID changes save ", {}), ("immediately", {"bold": True}),
              (" when you click ", {}), ("Save Product ID", {"bold": True}),
              (". They do ", {}), ("not", {"bold": True}),
              (" wait for the certificate's Save button.", {})]],
            [[("Certificate → Create", {"bold": True})],
             [("Product IDs are saved together with the certificate when you save the new record.",
               {})]],
        ],
        widths=[1.7, 4.8],
    )

    body(doc, "", space_after=4)
    body(doc, segments=[
        ("Practical guidance for analysts: ", {"bold": True}),
        ("on the Edit page, complete all Product ID additions, edits and removals as a group "
         "before leaving the page. Navigating away does not discard them — they are already saved.",
         {}),
    ])

    h2(doc, "Supported identifier types and expected formats")

    make_table(
        doc,
        ["Type", "Description", "Expected format"],
        [
            ["ASIN", "Amazon Standard Identification Number", "10 characters, letters and digits"],
            ["UPC", "Universal Product Code", "12 digits"],
            ["GTIN", "Global Trade Item Number", "8, 12, 13 or 14 digits"],
            ["EAN", "European Article Number", "13 digits"],
            ["ISBN", "International Standard Book Number", "ISBN-10 or ISBN-13"],
            ["SKU", "Stock Keeping Unit", "3–50 characters; letters, digits, hyphen, underscore"],
            ["MPN", "Manufacturer Part Number",
             "1–50 characters; letters, digits, hyphen, underscore, period"],
        ],
        widths=[0.85, 3.05, 2.6],
    )

    body(doc, "", space_after=4)
    body(doc, "Values that do not match the expected format are rejected at save with a message "
              "naming the type.")
    body(doc, segments=[
        ("Status. ", {"bold": True}),
        ("A Product ID's status is derived automatically from its parent certificate. It is not "
         "set by the user.", {}),
    ])

    # ================================================ Time Out
    h1(doc, "Time Out")

    body(doc, segments=[("30 minutes of inactivity.", {"bold": True, "color": GREEN_DARK})], size=12)

    bullet(doc, [
        ("The timeout is a ", {}), ("sliding", {"bold": True}),
        (" window — the 30-minute clock resets on every page load or action. It is 30 minutes of ",
         {}),
        ("inactivity", {"italic": True}),
        (", not 30 minutes per login.", {}),
    ])
    bullet(doc, [("Applies to both the login session and the sign-in cookie.", {})])

    body(doc, "", space_after=2)
    callout(doc, [
        ("What analysts need to know.  ", {"bold": True}),
        ("There is no on-screen warning before the session ends, and in-progress form data is not "
         "preserved. A certificate Create or Edit form left open for more than 30 minutes without "
         "activity will return the user to the login page on submit, and unsaved entries on that "
         "form are lost.", {}),
    ])

    body(doc, segments=[
        ("Recommended practice: ", {"bold": True}),
        ("save long certificate entries incrementally rather than composing the entire record in "
         "one sitting.", {}),
    ])

    # ================================================ Character Limits
    h1(doc, "Character Limits")

    body(doc, "Values exceeding these limits are rejected at save.")

    make_table(
        doc,
        ["Field", "Limit"],
        [
            ["Certificate Number", "50"],
            ["Status", "50"],
            [[("Location ", {}), ("(certificate and company)", {"italic": True})], "175"],
            ["Country", "100"],
            ["Partner", "100"],
            [[("Certification Results", {"bold": True})], [("1,000", {"bold": True})]],
            ["Company name", "255"],
            ["Product name", "255"],
            ["Product Category / Product Type / Certification Type name", "255"],
            ["Product ID value", "50"],
            ["Product ID — product name", "1,000"],
            ["Product ID — SKU", "100"],
            ["Product ID — description", "500"],
            ["Tag name", "50"],
            ["Contact info", "255"],
            ["Email address", "255"],
            ["Website", "500"],
        ],
        widths=[4.9, 1.6],
    )

    body(doc, "", space_after=4)
    body(doc, segments=[
        ("Note on error messaging. ", {"bold": True}),
        ("Certification Results returns a specific message on import (“exceeds maximum length "
         "of 1000 characters”). Other fields currently return a generic validation error "
         "rather than naming the limit. Flagged as a usability improvement; not a functional "
         "defect.", {}),
    ])

    # ================================================ Import Files
    h1(doc, "Import Files")

    bullet(doc, [
        ("CSV only. ", {"bold": True}),
        ("Excel (", {}), (".xlsx", {"mono": True}), ("), JSON and PDF are not accepted for import.",
                          {}),
    ])
    bullet(doc, [
        ("Always start from the template. ", {"bold": True}),
        ("Download the Template from the relevant page or section for the type of content being "
         "imported, populate it, and upload it. Column headers must match the template.", {}),
    ])
    bullet(doc, [("Maximum file size: 10 MB.", {"bold": True})])

    h2(doc, "Multiple entries in a single cell")

    body(doc, segments=[
        ("Separate items with a ", {}),
        ("comma", {"bold": True}),
        (" — not a semicolon, pipe, or any other character.", {}),
    ])

    code_block(doc, [
        "Tags:            recycled, post-consumer, verified",
        "Internal Tags:   priority-review, amazon-scope",
    ])

    body(doc, "Leading and trailing spaces around each item are trimmed automatically.")
    body(doc, segments=[
        ("Exception: the Partner column also accepts a pipe ( | ) as a separator. Commas work "
         "everywhere and are the recommended convention.", {"italic": True, "color": MUTED}),
    ], size=9.5)

    # ================================================ Change log
    h1(doc, "Change Log")

    make_table(
        doc,
        ["Date", "Change"],
        [
            ["14 Aug 2026",
             "Time Out and Character Limits answered from the application. Product Identifiers "
             "clarified — save behaviour differs between Create and Edit. Certificate Number "
             "five-digit rule and save-time assignment confirmed as previously written."],
        ],
        widths=[1.2, 5.3],
        zebra=False,
    )

    doc.core_properties.title = "Key System Backend Functions"
    doc.core_properties.subject = "GCC Database Solution"
    doc.core_properties.author = "Anuki Consulting"
    doc.core_properties.comments = "Version 1.0 — 14 August 2026"

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "Key_System_Backend_Functions.docx")
