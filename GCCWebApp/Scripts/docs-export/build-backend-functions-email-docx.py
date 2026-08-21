"""
Build the Word version of the client email: "Key System Backend Functions — open items answered".

Source of truth for the wording:
    docs/tracking/2026-08-14-backend-functions-email.md

Laid out as a review-ready email draft: a header block carrying Subject / To / From / Date, a rule,
then the body exactly as it should read in the client's inbox. The body is plain enough to paste
straight into Outlook — the tables survive a paste, the callouts flatten to indented text.

Content is authored here, not parsed from the markdown. Edit both together.

Usage:  python build-backend-functions-email-docx.py <output.docx>
"""
import sys
from docx import Document
from docx.shared import Pt, Inches

from docx_kit import (
    GREEN_DARK, GREEN_MID, MUTED, INK, QUOTE_FILL, QUOTE_RULE,
    base_document, body, bullet, h1, hrule, make_table, callout, rich, style_run,
)


def field_row(doc, label, value, bold_value=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.95))
    r = p.add_run(label)
    style_run(r, size=9.5, bold=True, color=MUTED)
    p.add_run("\t")
    style_run(p.add_run(value), size=10.5, bold=bold_value, color=INK)
    return p


def numbered_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    style_run(p.add_run(text), size=12, bold=True, color=GREEN_DARK)
    return p


def build(out_path):
    doc = Document()
    base_document(doc, "Email draft  ·  Key System Backend Functions  ·  14 August 2026")

    # ---------------------------------------------------------- header block
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    style_run(p.add_run("EMAIL DRAFT — FOR REVIEW"), size=9, bold=True, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    style_run(
        p.add_run("Key System Backend Functions — open items answered"),
        size=19, bold=True, color=GREEN_DARK,
    )

    field_row(doc, "SUBJECT", "Key System Backend Functions — open items answered "
                              "(timeout, character limits)", bold_value=True)
    field_row(doc, "TO", "[Client contact]")
    field_row(doc, "FROM", "Shawn West, Anuki Consulting")
    field_row(doc, "DATE", "14 August 2026")
    field_row(doc, "ATTACHMENT", "GCC Database - Key System Backend Functions v1.0.docx")

    hrule(doc, color="1B5E3F", size=12, space_before=10, space_after=12)

    # ---------------------------------------------------------- body
    body(doc, "Hi [Name],")

    body(doc, "Following up on the Key System Backend Functions section — the two outstanding "
              "questions are answered below, and I've confirmed the remaining items against the "
              "application itself rather than against earlier documentation. Two points needed a "
              "correction, noted at the end.")

    # -- 1
    numbered_heading(doc, "1.  Session Timeout")

    body(doc, segments=[
        ("The system times out after ", {}),
        ("30 minutes of inactivity", {"bold": True}),
        (". It is a sliding window — the 30-minute clock resets on every page load or action, so "
         "it's 30 minutes of inactivity rather than 30 minutes per login.", {}),
    ])

    body(doc, segments=[
        ("One point worth including in analyst-facing material: there is ", {}),
        ("no on-screen warning before the session ends, and in-progress form data is not preserved",
         {"bold": True}),
        (". A certificate Create or Edit form left open past the timeout will return the user to "
         "the login page on submit, and unsaved entries on that form are lost. We'd recommend the "
         "guide advise saving longer certificate entries incrementally.", {}),
    ])

    body(doc, "If a longer window is preferable, this is a straightforward configuration change — "
              "happy to set it to whatever period works for your team.")

    # -- 2
    numbered_heading(doc, "2.  Character Limits")

    make_table(
        doc,
        ["Field", "Limit"],
        [
            ["Certificate Number", "50"],
            ["Status", "50"],
            ["Location (certificate and company)", "175"],
            ["Country", "100"],
            ["Partner", "100"],
            [[("Certification Results", {"bold": True})], [("1,000", {"bold": True})]],
            ["Company name", "255"],
            ["Product name", "255"],
            ["Product Category / Type / Certification Type", "255"],
            ["Product ID value", "50"],
            ["Product ID — product name", "1,000"],
            ["Product ID — SKU", "100"],
            ["Product ID — description", "500"],
            ["Tag name", "50"],
            ["Contact info / Email", "255"],
            ["Website", "500"],
        ],
        widths=[4.9, 1.6],
    )

    body(doc, "", space_after=4)
    body(doc, segments=[
        ("Values over the limit are rejected at save. Imported files additionally have a ", {}),
        ("10 MB maximum file size", {"bold": True}),
        (".", {}),
    ])

    # -- 3
    numbered_heading(doc, "3.  Confirmed as documented — no change needed")

    body(doc, segments=[
        ("Certificate Number. ", {"bold": True, "italic": True}),
        ("Format ", {}), ("YY-####", {"mono": True}),
        (" (e.g. ", {}), ("26-0001", {"mono": True}),
        ("), auto-assigned. The sequence restarts at ", {}), ("0001", {"mono": True}),
        (" each calendar year — the first certificate of 2027 will be ", {}),
        ("27-0001", {"mono": True}),
        (", with nothing carried over. It widens to five digits only after all 9,999 numbers for "
         "a year are genuinely used; isolated high-numbered records from an import will not push "
         "new certificates into five digits.", {}),
    ])

    body(doc, segments=[
        ("The number is allocated ", {}),
        ("at the moment of save", {"bold": True}),
        (", not when the form is opened. If two analysts are working simultaneously, the first to "
         "save keeps the number shown to them and the second is automatically given the next "
         "available number. The same certificate number is never issued twice.", {}),
    ])

    body(doc, segments=[
        ("Import — multiple entries in one cell. ", {"bold": True, "italic": True}),
        ("Confirmed: separate items with a ", {}),
        ("comma", {"bold": True}),
        (", not a semicolon. For example: ", {}),
        ("recycled, post-consumer, verified", {"mono": True}),
        (". Spaces around each item are trimmed automatically. CSV only, always starting from the "
         "downloadable template.", {}),
    ])

    # -- 4
    numbered_heading(doc, "4.  Two corrections to the current draft")

    body(doc, segments=[
        ("Product Identifiers. ", {"bold": True, "italic": True}),
        ("The current wording — “auto saves updates upon revision, not on save” — is accurate on "
         "the ", {}),
        ("Edit", {"bold": True}),
        (" page only. There, Product ID changes save immediately when you click ", {}),
        ("Save Product ID", {"italic": True}),
        (", independent of the certificate's Save button. On the ", {}),
        ("Create", {"bold": True}),
        (" page, Product IDs are saved together with the new certificate. As written, the "
         "instruction reads as incorrect on one of the two pages, so we'd suggest splitting it:",
         {}),
    ])

    callout(doc, [
        ("On the Edit page, Product ID changes save immediately when you click Save Product ID — "
         "they do not wait for the certificate Save button. On the Create page, Product IDs are "
         "saved with the certificate.", {"italic": True}),
    ], fill=QUOTE_FILL, rule=QUOTE_RULE)

    body(doc, segments=[
        ("Migrated certificate numbers. ", {"bold": True, "italic": True}),
        ("The section states a certificate number is never assigned twice. That holds for every "
         "certificate created in the system. It's worth adding a footnote that a limited set of "
         "records migrated from the prior system do share a number, because those numbers were "
         "issued to more than one certificate before this system existed and cannot be changed "
         "after issue. Those records are individually flagged, and new certificates can never "
         "take an existing number.", {}),
    ])

    # -- close
    hrule(doc, space_before=12, space_after=10)

    body(doc, "I've attached the fully updated section with these answers incorporated. Let me "
              "know if you'd like the timeout adjusted or the Product Identifiers wording revised "
              "differently.")

    body(doc, "Best regards,", space_after=2)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    style_run(p.add_run("Shawn West"), size=10.5, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    style_run(p.add_run("Anuki Consulting"), size=10.5, color=MUTED)

    doc.core_properties.title = "Key System Backend Functions — open items answered"
    doc.core_properties.subject = "GCC Database Solution — client email draft"
    doc.core_properties.author = "Anuki Consulting"
    doc.core_properties.comments = "Email draft — 14 August 2026"

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "Backend_Functions_Email.docx")
