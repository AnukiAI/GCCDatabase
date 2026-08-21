"""
Shared python-docx formatting helpers for GCC client-facing Word deliverables.

Extracted from build-backend-functions-docx.py so multiple generators share one look:
dark-green headings, shaded table headers with zebra rows, accent-bar callouts, monospace
code blocks, page-number footers.

NOTE: build-backend-functions-docx.py predates this module and still carries its own copy of
these helpers. It is verified and shipped, so it was deliberately left alone. New generators
should import from here.
"""
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- palette
GREEN_DARK = RGBColor(0x1B, 0x5E, 0x3F)
GREEN_MID = RGBColor(0x2E, 0x7D, 0x54)
INK = RGBColor(0x21, 0x21, 0x21)
MUTED = RGBColor(0x5F, 0x6B, 0x66)
HDR_FILL = "1B5E3F"
ALT_FILL = "F2F7F4"
NOTE_FILL = "FFF8E1"
NOTE_RULE = "E0A800"
QUOTE_FILL = "F2F7F4"
QUOTE_RULE = "2E7D54"


# ---------------------------------------------------------------- low-level
def shade(element, hex_fill):
    pr = element._tc.get_or_add_tcPr() if hasattr(element, "_tc") else element._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def cell_borders(cell, color="D5E0DA", size=4, edges=("top", "bottom", "left", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def left_accent(cell, color, size=18):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:color"), color)
    borders.append(el)
    tcPr.append(borders)


def cell_margins(cell, top=80, bottom=80, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def hrule(doc, color="D5E0DA", size=6, space_before=6, space_after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), color)
    borders.append(el)
    pPr.append(borders)
    return p


def page_number_footer(section, left_text):
    p = section.footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    usable = section.page_width - section.left_margin - section.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(usable)

    r = p.add_run(left_text + "\t")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    r = p.add_run("Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    for instr in ("PAGE", "NUMPAGES"):
        run = p.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        itxt = OxmlElement("w:instrText")
        itxt.set(qn("xml:space"), "preserve")
        itxt.text = f" {instr} "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(itxt)
        run._r.append(end)
        if instr == "PAGE":
            sep = p.add_run(" of ")
            sep.font.size = Pt(8)
            sep.font.color.rgb = MUTED


# ---------------------------------------------------------------- text
def style_run(run, size=10.5, bold=False, color=INK, italic=False, mono=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    if mono:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def rich(paragraph, segments, size=10.5):
    """segments: list of (text, {bold, italic, mono, color})"""
    for text, opts in segments:
        r = paragraph.add_run(text)
        style_run(
            r,
            size=size,
            bold=opts.get("bold", False),
            italic=opts.get("italic", False),
            color=opts.get("color", INK),
            mono=opts.get("mono", False),
        )
    return paragraph


def body(doc, text=None, segments=None, size=10.5, space_after=8, indent=0, line_spacing=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    if segments:
        rich(p, segments, size=size)
    elif text is not None:
        style_run(p.add_run(text), size=size)
    return p


def bullet(doc, segments, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    rich(p, segments, size=size)
    return p


def h1(doc, text, size=15, space_before=20):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    style_run(p.add_run(text), size=size, bold=True, color=GREEN_DARK)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    style_run(p.add_run(text), size=11.5, bold=True, color=GREEN_MID)
    return p


# ---------------------------------------------------------------- blocks
def make_table(doc, headers, rows, widths=None, zebra=True, font_size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False

    hdr = t.rows[0]
    repeat_header(hdr)
    for i, htext in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        shade(c, HDR_FILL)
        cell_borders(c, color=HDR_FILL)
        cell_margins(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        style_run(p.add_run(htext), size=font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            c.text = ""
            cell_borders(c)
            cell_margins(c)
            if zebra and ri % 2 == 1:
                shade(c, ALT_FILL)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            segs = val if isinstance(val, list) else [(val, {})]
            rich(p, segs, size=font_size)

    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Inches(w)
    return t


def callout(doc, segments, fill=NOTE_FILL, rule=NOTE_RULE, width=6.5, size=10):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    c.text = ""
    shade(c, fill)
    left_accent(c, rule)
    cell_margins(c, top=130, bottom=130, left=170, right=140)
    c.width = Inches(width)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    rich(p, segments, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def code_block(doc, lines, width=6.5, size=9.5):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    c.text = ""
    shade(c, "F4F6F5")
    cell_borders(c, color="DDE4E0")
    cell_margins(c, top=120, bottom=120, left=150, right=140)
    c.width = Inches(width)
    first = True
    for line in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(2)
        style_run(p.add_run(line), size=size, mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def base_document(doc, footer_text, margins=(0.85, 0.85, 0.95, 0.95)):
    """Apply the house Normal style, margins and footer. Returns the section."""
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    s = doc.sections[0]
    top, bottom, left, right = margins
    s.top_margin = Inches(top)
    s.bottom_margin = Inches(bottom)
    s.left_margin = Inches(left)
    s.right_margin = Inches(right)
    page_number_footer(s, footer_text)
    return s
