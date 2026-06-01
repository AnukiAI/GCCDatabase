"""Convert client memo markdown to Word .docx.

Targeted parser for the client-memo-* family of files in docs/tracking/. Handles
headers, tables, bullet lists, paragraphs with inline bold/italic/code markdown,
and horizontal rules. Not a general markdown converter.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs_with_inline_formatting(paragraph, text):
    """Parse **bold**, *italic*, `code` inside a paragraph and add runs."""
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        elif token.startswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table_row(line):
    """Split a markdown table row '| a | b | c |' into cell strings."""
    cells = line.strip().strip('|').split('|')
    return [c.strip() for c in cells]


def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    # Base style for body text
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]

        # Skip empty lines (Word handles spacing between paragraphs)
        if not line.strip():
            i += 1
            continue

        # H1
        if line.startswith('# '):
            p = doc.add_heading(line[2:].strip(), level=0)
            i += 1
            continue

        # H2
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue

        # H3
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        # Horizontal rule
        if line.strip() == '---':
            p = doc.add_paragraph()
            run = p.add_run('─' * 60)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            i += 1
            continue

        # Table block: starts with '|' and the next line is the separator '|---|...'
        if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i + 1].strip()):
            header_cells = parse_table_row(line)
            # Skip separator row
            data_rows = []
            j = i + 2
            while j < len(lines) and lines[j].startswith('|'):
                data_rows.append(parse_table_row(lines[j]))
                j += 1

            table = doc.add_table(rows=1 + len(data_rows), cols=len(header_cells))
            table.style = 'Light Grid Accent 1'
            # Header row
            for col_idx, cell_text in enumerate(header_cells):
                cell = table.rows[0].cells[col_idx]
                cell.text = ''
                p = cell.paragraphs[0]
                run = p.add_run(cell_text)
                run.bold = True
            # Data rows
            for row_idx, row_data in enumerate(data_rows):
                for col_idx, cell_text in enumerate(row_data):
                    if col_idx >= len(header_cells):
                        continue
                    cell = table.rows[row_idx + 1].cells[col_idx]
                    cell.text = ''
                    add_runs_with_inline_formatting(cell.paragraphs[0], cell_text)
            doc.add_paragraph()  # spacer
            i = j
            continue

        # Bullet list
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs_with_inline_formatting(p, line[2:])
            i += 1
            continue

        # Numbered list
        m = re.match(r'^(\d+)\.\s+(.*)', line)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs_with_inline_formatting(p, m.group(2))
            i += 1
            continue

        # Paragraph (may continue across lines until blank)
        para_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#', '|', '- ', '---')) and not re.match(r'^\d+\.\s', lines[i]):
            para_lines.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        add_runs_with_inline_formatting(p, ' '.join(para_lines))

    doc.save(docx_path)
    print(f'Wrote: {docx_path}')


if __name__ == '__main__':
    md_in = sys.argv[1]
    docx_out = sys.argv[2] if len(sys.argv) > 2 else Path(md_in).with_suffix('.docx')
    convert(md_in, str(docx_out))
