# -*- coding: utf-8 -*-
"""
Renders a guide to BOTH formats (.docx + .html) from one markdown source.

    python build-guide.py                          # Accounts & Sign-In Guide (default)
    python build-guide.py ../partner-api-setup-guide.md

Reads   : the source .md  (the only file to edit)
          shots/*.png     (from capture-screenshots.js)
Writes  : ../<basename>.docx  and  ../<basename>.html

The output basename comes from the frontmatter `slug:` if present, otherwise from the source
filename. Passing no argument reproduces the original behaviour exactly, so the Accounts guide
still builds from ../guide.md to GCC_Accounts_SignIn_Guide_v1.0.*.

The two formats previously had separate copies of the text and drifted. Everything below is
rendering only — no wording lives here.
"""
import base64, glob, html as htmlmod, os, re, sys
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from guide_parser import parse

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "shots")
DEST = os.path.dirname(HERE)

# The source .md sits at the top of docs/user-guide/, beside the deliverables it produces —
# it is the file people edit, so it does not belong buried in _source/ with the machinery.
# Default (no argv) = the Accounts & Sign-In Guide, unchanged.
SOURCE_MD = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(DEST, "guide.md")
if not os.path.exists(SOURCE_MD):
    raise SystemExit(f"  Source not found: {SOURCE_MD}")
META, CONTENT = parse(SOURCE_MD)

# Output basename: frontmatter `slug:` wins, else the source filename. The Accounts guide has no
# slug and is named guide.md, so it keeps its historical output names via the explicit fallback.
#
# The slug is read from the raw file rather than from META: parse() returns a CURATED meta dict
# (title, lede, org, colophon, rows, toc, intro_callout) and folds everything else into `rows`,
# so META.get("slug") is always None. Read it here instead of widening the parser's contract.
def _slug_from(path):
    with open(path, encoding="utf-8") as fh:
        if fh.readline().strip() != "---":
            return None
        for line in fh:
            if line.strip() == "---":
                return None
            if line.split(":", 1)[0].strip() == "slug" and ":" in line:
                return line.split(":", 1)[1].strip() or None
    return None

_stem = os.path.splitext(os.path.basename(SOURCE_MD))[0]
BASENAME = (_slug_from(SOURCE_MD)
            or ("GCC_Accounts_SignIn_Guide_v1.0" if _stem == "guide" else _stem))
DOCX = os.path.join(DEST, BASENAME + ".docx")
HTML = os.path.join(DEST, BASENAME + ".html")

GREEN, NAVY, GREY = RGBColor(0x44, 0x76, 0x38), RGBColor(0x00, 0x2B, 0x76), RGBColor(0x4A, 0x5B, 0x50)
MAX_W_IN, MAX_H_IN = 6.0, 7.2
CALLOUT = {                       # kind -> (word accent, word fill, css class)
    "note": ("92C83E", "F2F8E6", "note"),
    "warn": ("8A5410", "FDF6EA", "note warn"),
    "stop": ("9C2B2B", "FCF0EF", "note stop"),
}
LABEL_COLOUR = {"note": "2F5527", "warn": "8A5410", "stop": "9C2B2B"}


_ITALIC = re.compile(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])")


def spans(text):
    """Inline formatting -> [(text, bold, italic, code)].

    `code` is taken first and is verbatim; then **bold**; then *italic* inside what is left.
    Before 2026-09-25 only **bold** was understood, so every `code` span and *italic* phrase was
    printed with its markers — 31 lines of the Partner guide showed raw backticks in BOTH formats.
    An odd number of backticks is treated as literal text rather than guessed at.
    """
    text = str(text)
    out = []
    chunks = text.split("`") if text.count("`") % 2 == 0 else [text]
    for ci, chunk in enumerate(chunks):
        if ci % 2 == 1:
            if chunk: out.append((chunk, False, False, True))
            continue
        for bi, part in enumerate(chunk.split("**")):
            bold, pos = bi % 2 == 1, 0
            for m in _ITALIC.finditer(part):
                if m.start() > pos: out.append((part[pos:m.start()], bold, False, False))
                out.append((m.group(1), bold, True, False))
                pos = m.end()
            if pos < len(part): out.append((part[pos:], bold, False, False))
    return [s for s in out if s[0]]


def add_runs(p, text, size=None):
    for tx, bo, it, co in spans(text):
        r = p.add_run(tx); r.bold = bo; r.italic = it
        if co:
            r.font.name = "Consolas"; r.font.size = Pt((size or 11) - 1)
        elif size:
            r.font.size = Pt(size)


def compress():
    for f in sorted(glob.glob(os.path.join(SHOTS, "*.png"))):
        im = Image.open(f).convert("RGB")
        if im.width > 1200:
            im = im.resize((1200, round(im.height * 1200 / im.width)), Image.LANCZOS)
        im.save(f[:-4] + ".jpg", "JPEG", quality=86, optimize=True, progressive=True)
    js = glob.glob(os.path.join(SHOTS, "*.jpg"))
    print(f"  screenshots : {len(js)} compressed, "
          f"{sum(os.path.getsize(x) for x in js)//1024} KB total")


# ══════════════════════════════════════════════════════════════ WORD
def render_docx():
    doc = Document()
    s = doc.sections[0]
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1.0)
    nm = doc.styles["Normal"]; nm.font.name = "Calibri"; nm.font.size = Pt(11)
    nm.paragraph_format.space_after = Pt(8); nm.paragraph_format.line_spacing = 1.12
    for name, size, before in (("Heading 1", 17, 20), ("Heading 2", 13, 14)):
        st = doc.styles[name]; st.font.name = "Calibri"; st.font.size = Pt(size)
        st.font.bold = True; st.font.color.rgb = GREEN
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(6); st.paragraph_format.keep_with_next = True

    def shade(cell, fill):
        e = OxmlElement("w:shd"); e.set(qn("w:val"), "clear"); e.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(e)

    def edge(cell, colour):
        tcPr = cell._tc.get_or_add_tcPr(); bs = OxmlElement("w:tcBorders")
        b = OxmlElement("w:left"); b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "18"); b.set(qn("w:color"), colour)
        bs.append(b); tcPr.append(bs)

    def para(text="", size=None, bold=False, italic=False, colour=None, after=None):
        p = doc.add_paragraph(); r = p.add_run(text); r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if colour: r.font.color.rgb = colour
        if after is not None: p.paragraph_format.space_after = Pt(after)

    def rich(text):
        add_runs(doc.add_paragraph(), text)

    def code_block(text):
        t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT
        c = t.cell(0, 0); shade(c, "F2F4F1"); c.width = Inches(6.0)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        for n, line in enumerate(text.split("\n")):
            r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            if n < text.count("\n"): r.add_break()
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Each numbered list gets its own w:num over the List Number abstract definition, with a
    # startOverride. Sharing the style's single numId made Word count every list in the document as
    # one — the Accounts guide's steps ran 1..26 and "Getting help" opened at step 24.
    numbering = doc.part.numbering_part.element
    ln_numid = doc.styles["List Number"].element.pPr.numPr.numId.val
    ln_abs = next(n.abstractNumId.val for n in numbering.num_lst if n.numId == ln_numid)

    def new_list(start):
        num = numbering.add_num(ln_abs)
        ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride"); so.set(qn("w:val"), str(start))
        ov.append(so); num.append(ov)
        return num.numId

    def table(headers, rows, widths=None, style="Light Grid Accent 3"):
        t = doc.add_table(rows=1, cols=len(headers)); t.style = style
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]; c.text = ""
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                add_runs(cells[i].paragraphs[0], val, size=9.5)
        if widths:
            for row in t.rows:
                for i, wv in enumerate(widths):
                    row.cells[i].width = Inches(wv)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def callout(kind, label, text):
        accent, fill, _ = CALLOUT[kind]
        t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT
        c = t.cell(0, 0); shade(c, fill); edge(c, accent); c.width = Inches(6.0)
        p = c.paragraphs[0]; r = p.add_run(label.replace("`", "").upper())  # labels are small caps; no code styling
        r.bold = True; r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(LABEL_COLOUR[kind])
        p.paragraph_format.space_after = Pt(2)
        p2 = c.add_paragraph()
        add_runs(p2, text, size=10)
        p2.paragraph_format.space_after = Pt(2)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    fign = [0]

    def figure(name, caption):
        path = os.path.join(SHOTS, name + ".jpg")
        if not os.path.exists(path):
            print("  MISSING", name); return
        w, h = Image.open(path).size
        scale = min(MAX_W_IN / (w / 96.0), MAX_H_IN / (h / 96.0), 1.0)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Inches(min((w / 96.0) * scale, MAX_W_IN)))
        fign[0] += 1
        cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = cp.add_run(f"Figure {fign[0]}  ")
        r1.bold = True; r1.font.size = Pt(9); r1.font.color.rgb = GREEN
        r2 = cp.add_run(caption); r2.font.size = Pt(9); r2.font.color.rgb = GREY; r2.italic = True
        cp.paragraph_format.space_after = Pt(14)

    # cover
    para(META["org"], size=10.5, bold=True, colour=GREEN, after=2)
    para(META["title"], size=30, bold=True, colour=NAVY, after=8)
    para(META["lede"], size=13, colour=GREY, after=18)
    table(["", ""], [list(r) for r in META["rows"]], widths=[1.5, 4.5])
    k, lab, txt = META["intro_callout"]; callout(k, lab, txt)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # contents — a real field, so Word builds and updates it
    doc.add_heading("Contents", level=1)
    p = doc.add_paragraph(); r = p.add_run()
    for tag, attr, text in (("begin", None, None), (None, "instr", 'TOC \\o "1-2" \\h \\z \\u'),
                            ("separate", None, None),
                            (None, "text", "Right-click and choose Update Field to build the contents."),
                            ("end", None, None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag)
        elif attr == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text
        else:
            e = OxmlElement("w:t"); e.text = text
        r._r.append(e)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for kind, payload in CONTENT:
        if kind == "h1":       doc.add_heading(payload, level=1)
        elif kind == "h2":     doc.add_heading(payload, level=2)
        elif kind == "p":      rich(payload)
        elif kind == "table":  table(payload[0], payload[1], payload[2])
        elif kind == "callout":callout(*payload)
        elif kind == "fig":    figure(*payload)
        elif kind == "pagebreak": doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "code":   code_block(payload)
        elif kind == "steps":
            start, items = payload
            nid = new_list(start)
            for it in items:
                p = doc.add_paragraph(style="List Number")
                numPr = p._p.get_or_add_pPr().get_or_add_numPr()
                numPr.get_or_add_ilvl().val = 0
                numPr.get_or_add_numId().val = nid
                add_runs(p, it)
                p.paragraph_format.space_after = Pt(4)

    para("")
    para(META["colophon"], size=9, colour=GREY, italic=True)

    ftr = doc.sections[0].footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = ftr.add_run(f"{META['org']} — {META['title']}    |    ")
    r0.font.size = Pt(8); r0.font.color.rgb = GREY
    r = ftr.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text
        r._r.append(e)
    r.font.size = Pt(8); r.font.color.rgb = GREY

    doc.core_properties.title = META["title"]
    doc.core_properties.subject = META["org"]
    doc.core_properties.author = "Anuki Consulting"
    doc.save(DOCX)
    print(f"  Word        : {os.path.basename(DOCX)}  {os.path.getsize(DOCX)//1024} KB, "
          f"{fign[0]} figures")


# ══════════════════════════════════════════════════════════════ WEB
CSS = """
:root{--paper:#fbfcfa;--card:#fff;--ink:#14231a;--ink-soft:#4a5b50;--ink-faint:#74857a;
--rule:#d9e0d7;--rule-soft:#eef2ec;--accent:#447638;--accent-ink:#2f5527;--lime:#92C83E;
--lime-wash:#f2f8e6;--navy:#002B76;--warn:#8a5410;--warn-wash:#fdf6ea;--danger:#9c2b2b;
--danger-wash:#fcf0ef;--on-accent:#fff;--shadow:0 1px 2px rgba(20,35,26,.06),0 6px 20px rgba(20,35,26,.06)}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--paper:#10150f;--card:#171d16;
--ink:#e6ece3;--ink-soft:#a7b3a3;--ink-faint:#7d8a7a;--rule:#2b352a;--rule-soft:#212a20;
--accent:#a3cf6d;--accent-ink:#c2e295;--lime:#92C83E;--lime-wash:#1d2617;--navy:#9db8ea;
--warn:#e0aa62;--warn-wash:#262015;--danger:#e79191;--danger-wash:#2a1a1a;--on-accent:#10150f;
--shadow:0 1px 2px rgba(0,0,0,.4),0 6px 20px rgba(0,0,0,.35)}}
:root[data-theme="dark"]{--paper:#10150f;--card:#171d16;--ink:#e6ece3;--ink-soft:#a7b3a3;
--ink-faint:#7d8a7a;--rule:#2b352a;--rule-soft:#212a20;--accent:#a3cf6d;--accent-ink:#c2e295;
--lime:#92C83E;--lime-wash:#1d2617;--navy:#9db8ea;--warn:#e0aa62;--warn-wash:#262015;
--danger:#e79191;--danger-wash:#2a1a1a;--on-accent:#10150f;
--shadow:0 1px 2px rgba(0,0,0,.4),0 6px 20px rgba(0,0,0,.35)}
*{box-sizing:border-box}
body{margin:0;background:var(--paper);color:var(--ink);font-size:16px;line-height:1.65;
font-family:"Segoe UI",-apple-system,BlinkMacSystemFont,Roboto,"Helvetica Neue",Arial,sans-serif;
-webkit-font-smoothing:antialiased}
.wrap{max-width:58rem;margin:0 auto;padding:clamp(1.5rem,4vw,3.5rem) clamp(1rem,4vw,2.5rem) 5rem}
.cover{border-top:4px solid var(--accent);padding-top:1.6rem;margin-bottom:2.2rem}
.eyebrow{font-size:.74rem;font-weight:700;letter-spacing:.16em;text-transform:uppercase;
color:var(--accent);margin:0 0 .7rem}
h1{font-size:clamp(2rem,5vw,2.9rem);line-height:1.1;margin:0 0 .8rem;letter-spacing:-.022em;
text-wrap:balance;font-weight:700}
.lede{font-size:1.09rem;color:var(--ink-soft);max-width:58ch;margin:0 0 1.5rem}
.meta{display:flex;flex-wrap:wrap;gap:.6rem 1.8rem;border-top:1px solid var(--rule);
border-bottom:1px solid var(--rule);padding:.9rem 0;font-size:.87rem;margin:0}
.meta div{display:flex;gap:.5rem;align-items:baseline}
.meta dt{color:var(--ink-faint);text-transform:uppercase;letter-spacing:.08em;font-size:.71rem;font-weight:700}
.meta dd{margin:0;font-weight:600}
nav.toc{background:var(--card);border:1px solid var(--rule);border-radius:8px;padding:1.2rem 1.4rem;margin:2rem 0 2.5rem}
nav.toc h2{margin:0 0 .7rem;font-size:.76rem;letter-spacing:.14em;text-transform:uppercase;
color:var(--accent);border:none;padding:0}
nav.toc ol{margin:0;padding-left:1.3rem;columns:2;column-gap:2rem}
@media (max-width:38rem){nav.toc ol{columns:1}}
nav.toc li{margin:.24rem 0;break-inside:avoid}
nav.toc a{color:var(--ink);text-decoration:none;border-bottom:1px solid transparent}
nav.toc a:hover{border-bottom-color:var(--accent);color:var(--accent)}
h2{font-size:1.42rem;margin:3rem 0 1rem;padding-bottom:.5rem;border-bottom:2px solid var(--accent);
letter-spacing:-.012em;font-weight:700;scroll-margin-top:1rem}
h3{font-size:1.08rem;margin:1.9rem 0 .6rem;font-weight:700}
p{max-width:66ch}
ol.steps{counter-reset:s;list-style:none;padding:0;margin:1rem 0;max-width:66ch}
ol.steps>li{counter-increment:s;position:relative;padding-left:2.4rem;margin:.75rem 0}
ol.steps>li::before{content:counter(s);position:absolute;left:0;top:.05rem;width:1.65rem;height:1.65rem;
background:var(--accent);color:var(--on-accent);border-radius:50%;display:grid;place-items:center;
font-size:.82rem;font-weight:700;font-variant-numeric:tabular-nums}
figure{margin:1.5rem 0 2rem;border:1px solid var(--rule);border-radius:8px;overflow:hidden;
background:var(--card);box-shadow:var(--shadow);break-inside:avoid}
figure img{display:block;width:100%;height:auto;border-bottom:1px solid var(--rule)}
figcaption{padding:.65rem .95rem;font-size:.86rem;color:var(--ink-soft)}
figcaption b{color:var(--accent);font-variant-numeric:tabular-nums;margin-right:.4rem}
.scroll{overflow-x:auto;margin:1rem 0 1.5rem;border:1px solid var(--rule);border-radius:7px;background:var(--card)}
table{border-collapse:collapse;width:100%;font-size:.93rem}
th,td{text-align:left;padding:.6rem .8rem;border-bottom:1px solid var(--rule-soft);vertical-align:top}
thead th{background:var(--rule-soft);font-size:.71rem;text-transform:uppercase;letter-spacing:.09em;
color:var(--ink-soft);font-weight:700;white-space:nowrap}
tbody tr:last-child td{border-bottom:none}
.note{border-left:3px solid var(--lime);background:var(--lime-wash);padding:.8rem 1.05rem;
margin:1.2rem 0;border-radius:0 6px 6px 0;font-size:.94rem;max-width:66ch}
.note .lab{display:block;font-size:.68rem;font-weight:800;letter-spacing:.13em;text-transform:uppercase;
color:var(--accent-ink);margin-bottom:.25rem}
.note p{margin:0;max-width:none}
.note.warn{border-left-color:var(--warn);background:var(--warn-wash)}
.note.warn .lab{color:var(--warn)}
.note.stop{border-left-color:var(--danger);background:var(--danger-wash)}
.note.stop .lab{color:var(--danger)}
footer.end{margin-top:3.5rem;padding-top:1.3rem;border-top:1px solid var(--rule);
font-size:.85rem;color:var(--ink-faint);max-width:66ch}
a{color:var(--navy)}
code{font-family:Consolas,"SFMono-Regular",Menlo,monospace;font-size:.9em;background:var(--rule-soft);
padding:.05em .32em;border-radius:3px}
pre.code{background:var(--rule-soft);border:1px solid var(--rule);border-radius:6px;padding:.8rem 1rem;
overflow-x:auto;max-width:66ch;margin:1rem 0 1.4rem}
pre.code code{background:none;padding:0;font-size:.88rem}
.pb{display:none}
@media print{body{background:#fff;font-size:10.5pt}nav.toc{break-after:page}
h2{break-after:avoid}figure{break-inside:avoid;box-shadow:none}.note{break-inside:avoid}
.pb{display:block;break-after:page}}
@media (prefers-reduced-motion:reduce){*{transition:none!important}}
"""


def esc(t):
    return htmlmod.escape(str(t), quote=False)


def inline(text):
    out = []
    for t, b, i, c in spans(text):
        h = f"<code>{esc(t)}</code>" if c else esc(t)
        if i: h = f"<em>{h}</em>"
        if b: h = f"<strong>{h}</strong>"
        out.append(h)
    return "".join(out)


def slug(h):
    return re.sub(r"[^a-z0-9]+", "-", re.sub(r"^\d+\.\s*", "", h).lower()).strip("-")


def render_html():
    out, fign = [], [0]
    out.append(f"<title>{esc(META['title'])}</title>\n<style>{CSS}</style>\n<div class='wrap'>")
    out.append("<header class='cover'>")
    out.append(f"<p class='eyebrow'>{esc(META['org'])}</p><h1>{esc(META['title'])}</h1>")
    out.append(f"<p class='lede'>{esc(META['lede'])}</p><dl class='meta'>")
    for k, v in META["rows"]:
        out.append(f"<div><dt>{esc(k)}</dt><dd>{esc(v)}</dd></div>")
    out.append("</dl></header>")
    k, lab, txt = META["intro_callout"]
    out.append(f"<div class='{CALLOUT[k][2]}'><span class='lab'>{esc(lab)}</span>"
               f"<p>{inline(txt)}</p></div>")

    out.append("<nav class='toc'><h2>Contents</h2><ol>")
    for t in META["toc"]:
        out.append(f"<li><a href='#{slug(t)}'>{esc(t)}</a></li>")
    out.append("</ol></nav>")

    for kind, payload in CONTENT:
        if kind == "h1":
            body = re.sub(r"^(\d+)\.\s*", "", payload)
            num = re.match(r"^(\d+)\.", payload)
            pre = f"<span style='color:var(--accent)'>{num.group(1)}</span>&nbsp; " if num else ""
            out.append(f"<h2 id='{slug(payload)}'>{pre}{esc(body)}</h2>")
        elif kind == "h2":
            out.append(f"<h3>{esc(payload)}</h3>")
        elif kind == "p":
            out.append(f"<p>{inline(payload)}</p>")
        elif kind == "steps":
            start, items = payload
            reset = f" style='counter-reset:s {start - 1}'" if start != 1 else ""
            out.append(f"<ol class='steps'{reset}>" +
                       "".join(f"<li>{inline(i)}</li>" for i in items) + "</ol>")
        elif kind == "code":
            out.append(f"<pre class='code'><code>{esc(payload)}</code></pre>")
        elif kind == "callout":
            ck, lab, txt = payload
            out.append(f"<div class='{CALLOUT[ck][2]}'><span class='lab'>{esc(lab.replace('`', ''))}</span>"
                       f"<p>{inline(txt)}</p></div>")
        elif kind == "pagebreak":
            out.append("<div class='pb'></div>")
        elif kind == "table":
            headers, rows, _ = payload
            out.append("<div class='scroll'><table><thead><tr>" +
                       "".join(f"<th>{esc(h)}</th>" for h in headers) + "</tr></thead><tbody>")
            for row in rows:
                out.append("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in row) + "</tr>")
            out.append("</tbody></table></div>")
        elif kind == "fig":
            name, caption = payload
            path = os.path.join(SHOTS, name + ".jpg")
            if not os.path.exists(path):
                print("  MISSING", name); continue
            fign[0] += 1
            b64 = base64.b64encode(open(path, "rb").read()).decode()
            # Never show a shot larger than it was captured. `figure img` is width:100%, so a narrow
            # crop (the 374px filter sidebar, the 390px phone view) was stretched to the full ~870px
            # column — 2.3x its real size, and far bigger than every other screenshot (reported on the
            # Company guide, 2026-09-28). Shots are captured at 1x, so natural width = on-screen width.
            nat_w = Image.open(path).size[0]
            cap = f" style='max-width:{nat_w + 2}px;margin-left:auto;margin-right:auto'" if nat_w < 900 else ""
            out.append(f"<figure{cap}><img src='data:image/jpeg;base64,{b64}' "
                       f"alt='{esc(caption.split(chr(46))[0])}' loading='lazy'>"
                       f"<figcaption><b>Figure {fign[0]}</b>{esc(caption)}</figcaption></figure>")

    out.append(f"<footer class='end'><p>{esc(META['colophon'])}</p></footer></div>")
    open(HTML, "w", encoding="utf-8", newline="").write("\n".join(out))
    print(f"  Web         : {os.path.basename(HTML)}  {os.path.getsize(HTML)//1024} KB, "
          f"{fign[0]} figures")


if __name__ == "__main__":
    compress()
    render_docx()
    render_html()
    # The .docx contents table is a Word field and stays a "Right-click and choose Update Field"
    # placeholder until Word computes it — which is what the client would have seen on page 2 of
    # every guide (caught 2026-09-25). Only Word can paginate, so hand the file to Word now rather
    # than trusting someone to remember a second step.
    import subprocess
    fin = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finalize-docx.ps1")
    try:
        res = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", fin,
                              "-Files", os.path.basename(DOCX)], capture_output=True, text=True, timeout=300)
        print(res.stdout.rstrip() or res.stderr.rstrip())
        if res.returncode != 0:
            raise RuntimeError("finalize-docx.ps1 failed")
    except Exception as e:
        print(f"\n  !! WORD CONTENTS NOT FILLED ({e}). The .docx still shows a placeholder on page 2 —"
              f"\n  !! run finalize-docx.ps1 on a machine with Word before sending it.")
    print(f"\n  Both formats rendered from {os.path.basename(SOURCE_MD)} — edit that file only.")
